import tkinter as tk
from tkinter import filedialog, font as tkfont, simpledialog
import ttkbootstrap as ttk
from ttkbootstrap import Style
import os
import shutil
import sys
import queue
import threading
import platform
from datetime import datetime, timedelta

from application.use_case.product_use_case import ProductCase
from application.use_case.sale_use_case import SaleCase
from infrastucture.db.db_manager import DBManager
from presentation.views.widgets import (
    apply_titlebar_theme, center_window, show_popup_smooth,
    get_menu_font, DarkMenuBar, MD, make_scrollable,
    make_scrolled_treeview,
    get_business_info, save_business_info, make_inline_business_header,
)
from presentation.views.payment_view import PaymentView
from presentation.views.inventory_view import InventoryView


def is_admin():
    try:
        import ctypes
        if platform.system() == "Windows":
            return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return False
    return False


def _parse_float(s):
    try:
        return float(str(s).replace("$", "").replace(".", "").replace(",", ".").strip() or 0)
    except Exception:
        return 0.0


class MainView(tk.Tk):
    def __init__(self):
        super().__init__()
        try:
            self.state('zoomed')
        except Exception:
            self.geometry("1400x900")

        # === Tema (leído desde settings) ===
        db_init = None
        try:
            if getattr(sys, 'frozen', False):
                base = os.path.dirname(sys.executable)
            else:
                base = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                                     "..", "..", ".."))
            db_init_path = os.path.join(base, "data", "ventas.db")
            if os.path.exists(db_init_path):
                from infrastucture.db.db_manager import DBManager as _DBM
                db_init = _DBM(db_init_path)
        except Exception:
            db_init = None

        self.current_theme = 'darkly'
        if db_init:
            try:
                self.current_theme = db_init.get_setting("theme", "darkly") or "darkly"
            except Exception:
                pass
        self.style = Style(theme=self.current_theme)
        self.configure(bg=self.style.colors.bg)
        self.is_fullscreen = False
        self._setup_dark_green_style()

        # === Base de datos ===
        if getattr(sys, 'frozen', False):
            base_dir = os.path.dirname(sys.executable)
        else:
            base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                                     "..", "..", ".."))
        self.base_dir = base_dir
        self.db_path = os.path.join(base_dir, "data", "ventas.db")
        os.makedirs(os.path.dirname(self.db_path), exist_ok=True)

        self.db_manager = DBManager(self.db_path)
        self.product_use_case = ProductCase(self.db_manager)
        self.sale_use_case = SaleCase(self.db_manager)

        # === Bandeja del sistema ===
        self.tray_icon = None
        self.tray_thread = None
        self.tray_queue = queue.Queue()
        self._reset_check_timer = None
        self._backup_check_timer = None

        # === Fase 2: auto-reset al iniciar ===
        self._do_auto_reset(silent=True)
        # === Fase 4: auto-backup al iniciar ===
        self._do_auto_backup(silent=True)
        # === Fase 6: abrir/verificar caja del día ===
        try:
            self.sale_use_case.get_or_create_today_cash_session()
        except Exception:
            pass

        self._update_window_title()

        if not self._check_startup_password():
            self.destroy()
            return

        self.create_widgets()
        self.show_page("pagos")
        self._apply_font_size()

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(200, self._poll_tray_queue)
        self.after(3000, self._schedule_daily_reset_check)
        self.after(4000, self._schedule_backup_check)

        self.bind('<F2>', lambda e: self.inventory_view.add_product_popup()
                  if self.current_page == "inventario" else None)
        self.bind('<F11>', lambda e: self.toggle_fullscreen())
        self.after(200, lambda: apply_titlebar_theme(
            self, self.current_theme == 'darkly'))

    # ============================================================
    # AUTO-RESET DIARIO (Fase 2)
    # ============================================================
    def _do_auto_reset(self, silent=False):
        try:
            reinicio = self.sale_use_case.auto_reset_if_new_day()
            if reinicio and not silent:
                MD.show_info(
                    "🕛 Ha cambiado el día.\n\n"
                    "El contador de ventas se reinició automáticamente.\n"
                    "La próxima venta será #01.",
                    "Reinicio automático", parent=self)
            return reinicio
        except Exception:
            return False

    def _schedule_daily_reset_check(self):
        self._check_daily_reset()

    def _check_daily_reset(self):
        try:
            reinicio = self._do_auto_reset(silent=False)
            if reinicio:
                try:
                    self.inventory_view.load_products()
                except Exception:
                    pass
                try:
                    self.sale_use_case.get_or_create_today_cash_session()
                except Exception:
                    pass
        except Exception:
            pass
        try:
            self._reset_check_timer = self.after(60000, self._check_daily_reset)
        except Exception:
            pass

    # ============================================================
    # RESPALDO AUTOMÁTICO (Fase 4)
    # ============================================================
    def _get_backup_folder(self):
        try:
            custom = self.db_manager.get_setting("backup_folder", "") or ""
            if custom:
                return custom
        except Exception:
            pass
        return os.path.join(os.path.dirname(self.db_path), "backups")

    def _is_auto_backup_enabled(self):
        try:
            return self.db_manager.get_setting("auto_backup_enabled", "1") == "1"
        except Exception:
            return True

    def _get_last_backup_date(self):
        try:
            return self.db_manager.get_setting("last_backup_date", "") or ""
        except Exception:
            return ""

    def _set_last_backup_date(self, date_str):
        try:
            self.db_manager.set_setting("last_backup_date", date_str)
        except Exception:
            pass

    def _do_auto_backup(self, silent=False):
        try:
            if not self._is_auto_backup_enabled():
                return False
            hoy = datetime.now().strftime("%Y-%m-%d")
            if self._get_last_backup_date() == hoy:
                return False
            folder = self._get_backup_folder()
            os.makedirs(folder, exist_ok=True)
            destino = os.path.join(folder, f"ventas_{hoy}.db")
            self.db_manager.close_connection()
            shutil.copy2(self.db_path, destino)
            self._set_last_backup_date(hoy)
            self.db_manager.get_connection()
            if not silent:
                MD.show_info(
                    f"💾 Respaldo automático creado:\n\n{destino}",
                    "Respaldo completado", parent=self)
            return True
        except Exception as e:
            if not silent:
                try:
                    MD.show_error(f"Error al respaldar: {e}", "Error", parent=self)
                except Exception:
                    pass
            try:
                self.db_manager.get_connection()
            except Exception:
                pass
            return False

    def _schedule_backup_check(self):
        self._check_auto_backup()

    def _check_auto_backup(self):
        try:
            self._do_auto_backup(silent=False)
        except Exception:
            pass
        try:
            self._backup_check_timer = self.after(3600000, self._check_auto_backup)
        except Exception:
            pass

    # ============================================================
    # TÍTULO
    # ============================================================
    def _update_window_title(self):
        try:
            info = get_business_info(self.db_manager)
            tipo = (info.get("type") or "").strip()
            nombre = (info.get("name") or "").strip()
            base = "MiniPOS Portable v2.5"
            if tipo and nombre:
                self.title(f"{base} — {tipo}, {nombre}")
            elif nombre:
                self.title(f"{base} — {nombre}")
            elif tipo:
                self.title(f"{base} — {tipo}")
            else:
                self.title(base)
        except Exception:
            self.title("MiniPOS Portable v2.5")

    def _refresh_business_header(self):
        try:
            if hasattr(self, "business_header") and self.business_header:
                self.business_header["refresh"]()
        except Exception:
            pass

    # ============================================================
    # BANDEJA
    # ============================================================
    def _poll_tray_queue(self):
        try:
            while True:
                action = self.tray_queue.get_nowait()
                if action == "show":
                    self._show_from_tray()
                elif action == "quit":
                    self._real_quit()
                    return
        except queue.Empty:
            pass
        except Exception:
            pass
        try:
            self.after(300, self._poll_tray_queue)
        except Exception:
            pass

    def _on_close(self):
        to_tray = self.db_manager.get_setting("close_to_tray", "0") == "1"
        if to_tray:
            self._hide_to_tray()
            return
        r = MD.yesno(
            "¿Estás seguro que deseas salir de MiniPOS Portable?\n\n"
            "Tus datos quedan guardados automáticamente.",
            "Confirmar salida", parent=self)
        if r == "Yes":
            self._real_quit()

    def _real_quit(self):
        try:
            if self._reset_check_timer:
                self.after_cancel(self._reset_check_timer)
        except Exception:
            pass
        try:
            if self._backup_check_timer:
                self.after_cancel(self._backup_check_timer)
        except Exception:
            pass
        try:
            if self.tray_icon:
                self.tray_icon.stop()
        except Exception:
            pass
        try:
            self.db_manager.close_connection()
        except Exception:
            pass
        try:
            self.destroy()
        except Exception:
            pass
        sys.exit(0)

    def _hide_to_tray(self):
        try:
            self.withdraw()
        except Exception:
            pass
        if self.tray_icon is None:
            self._create_tray_icon()

    def _create_tray_icon(self):
        try:
            from PIL import Image, ImageDraw
            import pystray
        except ImportError:
            try:
                self.deiconify()
                self.iconify()
            except Exception:
                pass
            return

        img = Image.new("RGBA", (64, 64), (10, 77, 31, 255))
        draw = ImageDraw.Draw(img)
        draw.rectangle([4, 4, 60, 60], outline=(168, 230, 168, 255), width=3)
        draw.rectangle([18, 24, 46, 28], fill=(168, 230, 168, 255))
        draw.rectangle([20, 16, 26, 48], fill=(168, 230, 168, 255))
        draw.rectangle([38, 16, 44, 48], fill=(168, 230, 168, 255))
        draw.rectangle([20, 30, 44, 34], fill=(168, 230, 168, 255))

        def on_show(icon, item):
            self.tray_queue.put("show")

        def on_quit(icon, item):
            self.tray_queue.put("quit")

        menu = pystray.Menu(
            pystray.MenuItem("Mostrar MiniPOS", on_show, default=True),
            pystray.MenuItem("Salir", on_quit),
        )
        self.tray_icon = pystray.Icon("MiniPOS", img, "MiniPOS Portable", menu)

        def run_tray():
            try:
                self.tray_icon.run()
            except Exception:
                pass

        self.tray_thread = threading.Thread(target=run_tray, daemon=True)
        self.tray_thread.start()

    def _show_from_tray(self):
        try:
            if self.tray_icon:
                self.tray_icon.stop()
        except Exception:
            pass
        self.tray_icon = None

        try:
            self.deiconify()
            self.state('zoomed')
            self.lift()
            self.focus_force()
        except Exception:
            pass

    # ============================================================
    # ESTILOS
    # ============================================================
    def _setup_dark_green_style(self):
        try:
            self.style.configure(
                "DarkGreen.TButton",
                background="#0a4d1f", foreground="#c8e6c9",
                borderwidth=0, focuscolor="#0a4d1f",
                padding=8, font=("Arial", 11, "bold"))
            self.style.map(
                "DarkGreen.TButton",
                background=[("active", "#083d18"), ("pressed", "#062e12"),
                            ("disabled", "#555555")],
                foreground=[("active", "#ffffff"), ("disabled", "#999999")])
        except Exception:
            pass

    # ============================================================
    # CONTRASEÑA DE INICIO
    # ============================================================
    def _check_startup_password(self):
        pwd = self.db_manager.get_setting("startup_password", "")
        if not pwd:
            return True
        self.withdraw()
        pop = tk.Toplevel(self)
        pop.title("MiniPOS - Iniciar sesión")
        pop.geometry("380x230")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.protocol("WM_DELETE_WINDOW", lambda: self._cancel_login(pop))
        pop.withdraw()

        tk.Label(pop, text="🔒 Contraseña requerida",
                 font=("Arial", 14, "bold"),
                 bg=self.style.colors.bg, fg=self.style.colors.fg).pack(pady=15)
        v = tk.StringVar()
        e = ttk.Entry(pop, textvariable=v, show="•", width=25,
                      font=("Arial", 14), justify="center")
        e.pack(pady=10)
        resultado = {"ok": False}

        def verificar(ev=None):
            if v.get() == pwd:
                resultado["ok"] = True
                pop.destroy()
            else:
                MD.show_error("Contraseña incorrecta", "Error", parent=pop)
                v.set("")
            return "break"

        ttk.Button(pop, text="Ingresar", command=verificar,
                   style="DarkGreen.TButton").pack(pady=10)
        e.bind("<Return>", verificar)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass

        def set_focus():
            try:
                if pop.winfo_exists():
                    e.focus_set()
            except Exception:
                pass
        pop.after(50, set_focus)
        pop.after(250, set_focus)

        self.wait_window(pop)
        if resultado["ok"]:
            self.deiconify()
            return True
        return False

    def _cancel_login(self, pop):
        pop.destroy()

    def _change_password(self):
        actual = self.db_manager.get_setting("startup_password", "")
        pop = tk.Toplevel(self)
        pop.title("Cambiar contraseña")
        pop.geometry("400x400")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="🔐 Cambiar contraseña",
                 font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=15)
        tk.Label(pop, text="Contraseña actual:", bg=bg, fg=fg).pack()
        e1 = ttk.Entry(pop, show="•", width=25, font=("Arial", 12))
        e1.pack(pady=5)
        tk.Label(pop, text="Nueva contraseña:", bg=bg, fg=fg).pack()
        e2 = ttk.Entry(pop, show="•", width=25, font=("Arial", 12))
        e2.pack(pady=5)
        tk.Label(pop, text="Confirmar nueva:", bg=bg, fg=fg).pack()
        e3 = ttk.Entry(pop, show="•", width=25, font=("Arial", 12))
        e3.pack(pady=5)
        vaciar = tk.BooleanVar(value=False)
        ttk.Checkbutton(pop, text="Quitar contraseña (dejar vacío)",
                        variable=vaciar).pack(pady=8)

        def aplicar(ev=None):
            if vaciar.get():
                self.db_manager.set_setting("startup_password", "")
                MD.show_info("Contraseña eliminada.", "Listo", parent=pop)
                pop.destroy()
                return "break"
            if e1.get() != actual:
                MD.show_error("La contraseña actual no coincide.", "Error",
                              parent=pop)
                return "break"
            if len(e2.get()) < 4:
                MD.show_error(
                    "La nueva contraseña debe tener al menos 4 caracteres.",
                    "Error", parent=pop)
                return "break"
            if e2.get() != e3.get():
                MD.show_error("Las contraseñas nuevas no coinciden.", "Error",
                              parent=pop)
                return "break"
            self.db_manager.set_setting("startup_password", e2.get())
            MD.show_info("Contraseña cambiada.", "Listo", parent=pop)
            pop.destroy()
            return "break"

        ttk.Button(pop, text="Guardar", command=aplicar,
                   style="DarkGreen.TButton").pack(pady=12)
        for e in (e1, e2, e3):
            e.bind("<Return>", aplicar)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass
        pop.after(50, lambda: e1.focus_set() if pop.winfo_exists() else None)

    def _setup_password_first_time(self):
        pop = tk.Toplevel(self)
        pop.title("Configurar contraseña")
        pop.geometry("380x280")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="🔐 Crear contraseña de inicio",
                 font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=15)
        tk.Label(pop, text="Nueva contraseña (mín. 4 caracteres):",
                 bg=bg, fg=fg).pack()
        e1 = ttk.Entry(pop, show="•", width=25, font=("Arial", 12))
        e1.pack(pady=5)
        tk.Label(pop, text="Confirmar:", bg=bg, fg=fg).pack()
        e2 = ttk.Entry(pop, show="•", width=25, font=("Arial", 12))
        e2.pack(pady=5)

        def guardar(ev=None):
            if len(e1.get()) < 4:
                MD.show_error("Mínimo 4 caracteres.", "Error", parent=pop)
                return "break"
            if e1.get() != e2.get():
                MD.show_error("No coinciden.", "Error", parent=pop)
                return "break"
            self.db_manager.set_setting("startup_password", e1.get())
            MD.show_info("Contraseña activada. Se pedirá al iniciar.",
                         "Listo", parent=pop)
            pop.destroy()
            return "break"

        ttk.Button(pop, text="Guardar", command=guardar,
                   style="DarkGreen.TButton").pack(pady=12)
        e1.bind("<Return>", guardar)
        e2.bind("<Return>", guardar)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass
        pop.after(50, lambda: e1.focus_set() if pop.winfo_exists() else None)

    # ============================================================
    # DATOS DEL NEGOCIO
    # ============================================================
    def _open_business_popup(self):
        pop = tk.Toplevel(self)
        pop.title("🏪 Datos del negocio")
        pop.geometry("520x620")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        info = get_business_info(self.db_manager)
        container = tk.Frame(pop, bg=bg)
        container.pack(fill="both", expand=True)
        refs = {}

        def build_content(parent):
            tk.Label(parent, text="🏪 DATOS DEL NEGOCIO",
                     font=("Arial", 15, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
            tk.Label(parent,
                     text="Estos datos aparecen en la barra de título, en Pagos, "
                          "Inventario y tickets.",
                     font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))
            campos = [
                ("type", "Tipo de negocio (ej. tienda, víveres, supermercado):"),
                ("name", "Nombre del negocio (ej. el chontico):"),
                ("owner", "Dueño (ej. Rosa):"),
                ("place", "Lugar / Dirección:"),
                ("phone", "Teléfono:"),
                ("employees", "Empleados (separados por coma):"),
            ]
            for key, label in campos:
                ttk.Label(parent, text=label, font=("Arial", 10)).pack(
                    pady=(10, 3), padx=20, anchor="w")
                entry = ttk.Entry(parent, width=52, font=("Arial", 11))
                entry.pack(pady=3, padx=20)
                if info.get(key):
                    entry.insert(0, info.get(key))
                refs[key] = entry
            ttk.Label(parent, text="Notas adicionales:",
                      font=("Arial", 10)).pack(pady=(10, 3), padx=20, anchor="w")
            notas = tk.Text(parent, width=52, height=4, font=("Arial", 10),
                            bg=bg, fg=fg, insertbackground=fg,
                            relief="solid", borderwidth=1, wrap="word")
            notas.pack(pady=3, padx=20)
            if info.get("notes"):
                notas.insert("1.0", info.get("notes"))
            refs["notes"] = notas

        def build_bottom(parent):
            def guardar():
                data = {}
                for k in ("type", "name", "owner", "place", "phone", "employees"):
                    try:
                        data[k] = refs[k].get()
                    except Exception:
                        data[k] = ""
                try:
                    data["notes"] = refs["notes"].get("1.0", tk.END).strip()
                except Exception:
                    data["notes"] = ""
                try:
                    save_business_info(self.db_manager, data)
                    self._update_window_title()
                    self._refresh_business_header()
                    pop.destroy()
                    MD.show_info("✅ Datos del negocio guardados.", "Listo",
                                 parent=self)
                except Exception as e:
                    MD.show_error(f"Error al guardar: {e}", "Error", parent=pop)

            ttk.Button(parent, text="💾 Guardar", command=guardar,
                       style="DarkGreen.TButton").pack(side="right", padx=6)
            ttk.Button(parent, text="Cancelar",
                       command=pop.destroy).pack(side="right", padx=6)

        make_scrollable(container, build_content, build_bottom, bg=bg)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass

    # ============================================================
    # CONFIGURACIONES SIMPLES
    # ============================================================
    def _configure_default_margin(self):
        actual = self.db_manager.get_setting("default_margin", "20") or "20"
        pop = tk.Toplevel(self)
        pop.title("📈 Margen por defecto")
        pop.geometry("400x260")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg
        tk.Label(pop, text="📈 Margen de ganancia por defecto",
                 font=("Arial", 13, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
        tk.Label(pop, text="Este % se usará por defecto al crear productos nuevos.",
                 font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))
        tk.Label(pop, text="% Ganancia:", font=("Arial", 11), bg=bg, fg=fg).pack()
        v = tk.StringVar(value=actual)
        e = ttk.Entry(pop, textvariable=v, width=15, font=("Arial", 14),
                      justify="center")
        e.pack(pady=8)
        e.select_range(0, tk.END)

        def guardar(ev=None):
            try:
                val = float(v.get().replace(",", ".").strip() or "0")
                if val < 0:
                    val = 0
            except ValueError:
                MD.show_error("Valor inválido", "Error", parent=pop)
                return "break"
            self.db_manager.set_setting("default_margin", str(val))
            pop.destroy()
            MD.show_info(f"✅ Margen por defecto: {val:g}%", "Listo", parent=self)
            return "break"

        ttk.Button(pop, text="Guardar", command=guardar,
                   style="DarkGreen.TButton").pack(pady=12)
        e.bind("<Return>", guardar)
        show_popup_smooth(pop)
        pop.after(80, lambda: e.focus_set() if pop.winfo_exists() else None)

    def _configure_low_stock(self):
        actual = self.db_manager.get_setting("low_stock_threshold", "5") or "5"
        pop = tk.Toplevel(self)
        pop.title("⚠️ Alerta de stock bajo")
        pop.geometry("400x260")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg
        tk.Label(pop, text="⚠️ Alerta de stock bajo",
                 font=("Arial", 13, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
        tk.Label(pop,
                 text="Los productos con stock ≤ este valor se marcarán en rojo.",
                 font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))
        tk.Label(pop, text="Unidades mínimas:", font=("Arial", 11),
                 bg=bg, fg=fg).pack()
        v = tk.StringVar(value=actual)
        e = ttk.Entry(pop, textvariable=v, width=15, font=("Arial", 14),
                      justify="center")
        e.pack(pady=8)
        e.select_range(0, tk.END)

        def guardar(ev=None):
            try:
                val = int(float(v.get().strip() or "0"))
                if val < 0:
                    val = 0
            except ValueError:
                MD.show_error("Valor inválido", "Error", parent=pop)
                return "break"
            self.db_manager.set_setting("low_stock_threshold", str(val))
            pop.destroy()
            try:
                self.inventory_view.load_products()
            except Exception:
                pass
            MD.show_info(f"✅ Alerta configurada: stock ≤ {val}", "Listo",
                         parent=self)
            return "break"

        ttk.Button(pop, text="Guardar", command=guardar,
                   style="DarkGreen.TButton").pack(pady=12)
        e.bind("<Return>", guardar)
        show_popup_smooth(pop)
        pop.after(80, lambda: e.focus_set() if pop.winfo_exists() else None)

    def _configure_expiry_alerts(self):
        dias_warning = self.db_manager.get_setting("expiry_days_warning", "15") or "15"
        dias_soon = self.db_manager.get_setting("expiry_days_soon", "7") or "7"
        dias_offer = self.db_manager.get_setting("expiry_days_offer", "2") or "2"

        pop = tk.Toplevel(self)
        pop.title("📅 Alertas de vencimiento")
        pop.geometry("500x460")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        container = tk.Frame(pop, bg=bg)
        container.pack(fill="both", expand=True)
        refs = {}

        def build_content(parent):
            tk.Label(parent, text="📅 Alertas de vencimiento",
                     font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
            tk.Label(parent,
                     text="Configura cuántos días antes quieres ver las alertas.",
                     font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))

            tk.Label(parent, text="🟠 Alerta temprana (días antes):",
                     font=("Arial", 11, "bold"), bg=bg, fg=fg,
                     anchor="w").pack(fill="x", padx=25, pady=(10, 2))
            tk.Label(parent, text="Aviso amarillo: producto por vencer.",
                     font=("Arial", 9, "italic"), bg=bg, fg="#888888",
                     anchor="w").pack(fill="x", padx=25)
            v_warning = tk.StringVar(value=dias_warning)
            e_warning = ttk.Entry(parent, textvariable=v_warning, width=15,
                                   font=("Arial", 12), justify="center")
            e_warning.pack(pady=4, padx=25, anchor="w")
            refs["warning"] = v_warning

            tk.Label(parent, text="🟡 Alerta cercana (días antes):",
                     font=("Arial", 11, "bold"), bg=bg, fg=fg,
                     anchor="w").pack(fill="x", padx=25, pady=(10, 2))
            tk.Label(parent,
                     text="Aviso naranja más fuerte: producto por vencer pronto.",
                     font=("Arial", 9, "italic"), bg=bg, fg="#888888",
                     anchor="w").pack(fill="x", padx=25)
            v_soon = tk.StringVar(value=dias_soon)
            e_soon = ttk.Entry(parent, textvariable=v_soon, width=15,
                               font=("Arial", 12), justify="center")
            e_soon.pack(pady=4, padx=25, anchor="w")
            refs["soon"] = v_soon

            tk.Label(parent, text="💡 Zona de oferta (días antes):",
                     font=("Arial", 11, "bold"), bg=bg, fg=fg,
                     anchor="w").pack(fill="x", padx=25, pady=(10, 2))
            tk.Label(parent,
                     text="Cuando falten X días, aparecerá la sugerencia "
                          "de poner el producto en oferta.",
                     font=("Arial", 9, "italic"), bg=bg, fg="#888888",
                     anchor="w", wraplength=440,
                     justify="left").pack(fill="x", padx=25)
            v_offer = tk.StringVar(value=dias_offer)
            e_offer = ttk.Entry(parent, textvariable=v_offer, width=15,
                                font=("Arial", 12), justify="center")
            e_offer.pack(pady=4, padx=25, anchor="w")
            refs["offer"] = v_offer

        def build_bottom(parent):
            def guardar():
                try:
                    w = max(0, int(_parse_float(refs["warning"].get())))
                    s = max(0, int(_parse_float(refs["soon"].get())))
                    o = max(0, int(_parse_float(refs["offer"].get())))
                except Exception:
                    MD.show_error("Valores inválidos", "Error", parent=pop)
                    return
                if not (w >= s >= o):
                    MD.show_error(
                        "Los días deben ser: temprana ≥ cercana ≥ oferta.\n"
                        "Ejemplo: 15 ≥ 7 ≥ 2",
                        "Orden inválido", parent=pop)
                    return
                self.db_manager.set_setting("expiry_days_warning", str(w))
                self.db_manager.set_setting("expiry_days_soon", str(s))
                self.db_manager.set_setting("expiry_days_offer", str(o))
                # También guardar en las keys nuevas que usa product_use_case
                try:
                    self.db_manager.set_setting("expiry_warn_days_1", str(w))
                    self.db_manager.set_setting("expiry_warn_days_2", str(s))
                    self.db_manager.set_setting("expiry_offer_days", str(o))
                except Exception:
                    pass
                try:
                    self.inventory_view.load_products()
                except Exception:
                    pass
                pop.destroy()
                MD.show_info(
                    f"✅ Alertas configuradas:\n"
                    f"🟠 Temprana: {w} días\n"
                    f"🟡 Cercana: {s} días\n"
                    f"💡 Oferta: {o} días",
                    "Listo", parent=self)

            ttk.Button(parent, text="💾 Guardar", command=guardar,
                       style="DarkGreen.TButton").pack(side="right", padx=6)
            ttk.Button(parent, text="Cancelar",
                       command=pop.destroy).pack(side="right", padx=6)

        make_scrollable(container, build_content, build_bottom, bg=bg)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass
        pop.after(100, lambda: e_warning.focus_set() if pop.winfo_exists() else None)

    def _configure_backup(self):
        habilitado = self._is_auto_backup_enabled()
        pop = tk.Toplevel(self)
        pop.title("💾 Respaldo automático")
        pop.geometry("560x400")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="💾 Respaldo automático diario",
                 font=("Arial", 13, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
        tk.Label(pop,
                 text="Se crea una copia de la base de datos cada día.\n"
                      "Se verifica al iniciar la app y cada hora.",
                 font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8",
                 justify="center").pack(pady=(0, 10))

        var_enabled = tk.BooleanVar(value=habilitado)
        ttk.Checkbutton(pop, text="✅ Activar respaldo automático",
                        variable=var_enabled,
                        bootstyle="success-round-toggle").pack(pady=8)

        tk.Label(pop, text="Carpeta de respaldos:",
                 font=("Arial", 11), bg=bg, fg=fg).pack(pady=(10, 3))
        carpeta_var = tk.StringVar(value=self._get_backup_folder())
        entry_frame = tk.Frame(pop, bg=bg)
        entry_frame.pack(pady=4, fill="x", padx=20)
        ttk.Entry(entry_frame, textvariable=carpeta_var, width=45,
                  font=("Arial", 10)).pack(side="left", fill="x", expand=True)

        def elegir_carpeta():
            folder = filedialog.askdirectory(
                initialdir=carpeta_var.get() or self.base_dir, parent=pop)
            if folder:
                carpeta_var.set(folder)

        ttk.Button(entry_frame, text="📁", command=elegir_carpeta,
                   bootstyle="info", width=3).pack(side="left", padx=4)

        def guardar():
            self.db_manager.set_setting("auto_backup_enabled",
                                        "1" if var_enabled.get() else "0")
            self.db_manager.set_setting("backup_folder",
                                         carpeta_var.get().strip())
            pop.destroy()
            MD.show_info("✅ Configuración de respaldo guardada.",
                         "Listo", parent=self)

        def respaldar_ahora():
            folder = carpeta_var.get().strip()
            try:
                os.makedirs(folder, exist_ok=True)
                hoy = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                destino = os.path.join(folder, f"ventas_manual_{hoy}.db")
                self.db_manager.close_connection()
                shutil.copy2(self.db_path, destino)
                self.db_manager.get_connection()
                MD.show_info(f"✅ Respaldo creado:\n{destino}", "Listo",
                             parent=pop)
            except Exception as e:
                MD.show_error(f"Error: {e}", "Error", parent=pop)

        bf = tk.Frame(pop, bg=bg)
        bf.pack(side="bottom", pady=15)
        ttk.Button(bf, text="💾 Guardar", command=guardar,
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(bf, text="📥 Respaldar ahora",
                   command=respaldar_ahora, bootstyle="info").pack(
                       side="left", padx=5)
        ttk.Button(bf, text="Cancelar", command=pop.destroy).pack(
            side="left", padx=5)

        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass

    # ============================================================
    # EXPORTAR INVENTARIO A EXCEL
    # ============================================================
    def export_inventory_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            MD.show_error("Falta openpyxl.", "Error", parent=self)
            return

        hoy = datetime.now().strftime("%Y-%m-%d")
        arch = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=f"inventario_{hoy}.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not arch:
            return

        try:
            products = self.product_use_case.list_products()
        except Exception:
            products = []

        if not products:
            MD.show_info("No hay productos en el inventario.",
                         "Sin datos", parent=self)
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventario"
        headers = ["ID", "Nombre", "Grupo", "Código",
                   "Costo", "% Gan.", "Precio real", "Precio redondeado",
                   "Stock", "Vence", "Estado",
                   "Valor costo", "Valor venta"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(horizontal="center")

        total_costo = 0.0
        total_venta = 0.0
        for p in products:
            valor_costo = (getattr(p, "cost", 0) or 0) * (p.stock or 0)
            valor_venta = (getattr(p, "rounded_price", 0) or p.price or 0) * (p.stock or 0)
            total_costo += valor_costo
            total_venta += valor_venta
            estado = "Activo"
            if getattr(p, "paused", 0):
                estado = "Pausado"
            elif getattr(p, "expiry_date", ""):
                try:
                    fecha = datetime.strptime(p.expiry_date, "%Y-%m-%d").date()
                    dias = (fecha - datetime.now().date()).days
                    if dias < 0:
                        estado = f"Vencido ({abs(dias)}d)"
                    elif dias <= 7:
                        estado = f"Vence en {dias}d"
                except Exception:
                    pass
            ws.append([
                p.product_id, p.name, getattr(p, "group_name", "") or "-",
                p.barcode or "-",
                getattr(p, "cost", 0) or 0,
                getattr(p, "margin_percent", 20) or 20,
                p.price or 0,
                getattr(p, "rounded_price", 0) or 0,
                p.stock or 0,
                getattr(p, "expiry_date", "") or "-",
                estado, valor_costo, valor_venta
            ])

        ws.append([])
        ws.append(["", "", "", "", "", "", "", "", "", "", "TOTALES:",
                   total_costo, total_venta])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)

        for col, w in zip("ABCDEFGHIJKLM",
                          [8, 30, 15, 18, 10, 8, 12, 14, 10, 12, 15, 14, 14]):
            ws.column_dimensions[col].width = w

        try:
            wb.save(arch)
            MD.show_info(
                f"✅ Inventario exportado:\n\n{arch}\n\n"
                f"Productos: {len(products)}\n"
                f"Valor costo total: ${total_costo:,.0f}\n"
                f"Valor venta total: ${total_venta:,.0f}".replace(",", "."),
                "Éxito", parent=self)
        except Exception as e:
            MD.show_error(f"Error al guardar: {e}", "Error", parent=self)

    # ============================================================
    # EXPORTAR HISTORIAL DE PRECIOS (FASE 5)
    # ============================================================
    def export_price_history(self):
        pop = tk.Toplevel(self)
        pop.title("📜 Exportar Historial de Precios")
        pop.geometry("480x400")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="📜 Exportar historial de precios",
                 font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
        tk.Label(pop, text="Elige el formato y el alcance del reporte.",
                 font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))

        tk.Label(pop, text="Formato:", font=("Arial", 11, "bold"),
                 bg=bg, fg=fg).pack(anchor="w", padx=25, pady=(10, 4))
        formato_var = tk.StringVar(value="TXT")
        ff = tk.Frame(pop, bg=bg)
        ff.pack(anchor="w", padx=25)
        for fmt in ("TXT", "Excel", "PDF"):
            ttk.Radiobutton(ff, text=fmt, variable=formato_var, value=fmt,
                            bootstyle="info").pack(side="left", padx=8)

        tk.Label(pop, text="Alcance:", font=("Arial", 11, "bold"),
                 bg=bg, fg=fg).pack(anchor="w", padx=25, pady=(15, 4))
        alcance_var = tk.StringVar(value="all")
        af = tk.Frame(pop, bg=bg)
        af.pack(anchor="w", padx=25)
        ttk.Radiobutton(af, text="Todos los productos",
                        variable=alcance_var, value="all",
                        bootstyle="info").pack(anchor="w")
        ttk.Radiobutton(af, text="Un producto específico",
                        variable=alcance_var, value="one",
                        bootstyle="info").pack(anchor="w", pady=3)

        prod_frame = tk.Frame(pop, bg=bg)
        prod_frame.pack(fill="x", padx=25, pady=(8, 4))
        tk.Label(prod_frame, text="Producto:", font=("Arial", 10),
                 bg=bg, fg=fg).pack(side="left", padx=(0, 5))
        try:
            products = self.product_use_case.list_products()
        except Exception:
            products = []
        product_map = {}
        labels = []
        for p in products:
            lbl = f"{p.name}  |  {p.barcode}" if p.barcode else p.name
            product_map[lbl] = p
            labels.append(lbl)
        prod_var = tk.StringVar()
        prod_combo = ttk.Combobox(prod_frame, textvariable=prod_var,
                                   width=42, state="disabled",
                                   values=labels)
        prod_combo.pack(side="left", fill="x", expand=True)

        def on_alcance_change(*args):
            if alcance_var.get() == "one":
                prod_combo.configure(state="readonly")
            else:
                prod_combo.configure(state="disabled")

        alcance_var.trace_add("write", on_alcance_change)

        def exportar():
            fmt = formato_var.get()
            alcance = alcance_var.get()
            if alcance == "one":
                sel = prod_var.get().strip()
                if not sel or sel not in product_map:
                    MD.show_error("Selecciona un producto de la lista.",
                                  "Error", parent=pop)
                    return
                prod = product_map[sel]
                hist = self.product_use_case.get_price_history(prod.product_id)
                if not hist:
                    MD.show_info("Ese producto no tiene historial.",
                                 "Sin datos", parent=pop)
                    return
                for h in hist:
                    h["product_name"] = prod.name
                    h["product_barcode"] = prod.barcode
            else:
                hist = self.product_use_case.get_all_price_history()
                if not hist:
                    MD.show_info("No hay historial de precios aún.",
                                 "Sin datos", parent=pop)
                    return
            pop.destroy()
            if fmt == "TXT":
                self._export_history_txt(hist)
            elif fmt == "Excel":
                self._export_history_excel(hist)
            else:
                self._export_history_pdf(hist)

        bf = tk.Frame(pop, bg=bg)
        bf.pack(side="bottom", pady=15)
        ttk.Button(bf, text="📤 Exportar", command=exportar,
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(bf, text="Cancelar",
                   command=pop.destroy).pack(side="left", padx=5)

        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass

    def _export_history_txt(self, hist):
        hoy = datetime.now().strftime("%Y-%m-%d")
        arch = filedialog.asksaveasfilename(
            defaultextension=".txt",
            initialfile=f"historial_precios_{hoy}.txt",
            filetypes=[("Texto", "*.txt")])
        if not arch:
            return
        L = ["=" * 90,
             "      HISTORIAL DE PRECIOS",
             f"      Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
             f"      Cambios: {len(hist)}",
             "=" * 90, ""]
        for h in hist:
            dif = h["new_price"] - h["old_price"]
            signo = "+" if dif > 0 else ""
            L.append(
                f"📅 {h['date']}  |  "
                f"{h.get('product_name','(producto eliminado)')[:35]:<35}  "
                f"|  Antes: ${h['old_price']:>10,.0f}  ->  "
                f"Ahora: ${h['new_price']:>10,.0f}  "
                f"({signo}${dif:>10,.0f})".replace(",", "."))
        L.append("")
        L.append("=" * 90)
        try:
            with open(arch, "w", encoding="utf-8") as f:
                f.write("\n".join(L))
            MD.show_info(f"Guardado:\n{arch}", "Listo", parent=self)
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def _export_history_excel(self, hist):
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            MD.show_error("Falta openpyxl.", "Error", parent=self)
            return
        hoy = datetime.now().strftime("%Y-%m-%d")
        arch = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=f"historial_precios_{hoy}.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not arch:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial"
        headers = ["Fecha", "Producto", "Código", "Precio antes",
                   "Precio ahora", "Diferencia"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(horizontal="center")
        for h in hist:
            dif = h["new_price"] - h["old_price"]
            ws.append([
                h["date"],
                h.get("product_name", "(producto eliminado)"),
                h.get("product_barcode", ""),
                h["old_price"], h["new_price"], dif
            ])
        for col, w in zip("ABCDEF", [20, 32, 18, 14, 14, 14]):
            ws.column_dimensions[col].width = w
        try:
            wb.save(arch)
            MD.show_info(f"Guardado:\n{arch}", "Listo", parent=self)
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def _export_history_pdf(self, hist):
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib import colors
            from reportlab.lib.units import cm
            from reportlab.platypus import (SimpleDocTemplate, Table,
                                             TableStyle, Paragraph, Spacer)
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError:
            MD.show_error(
                "Para exportar a PDF necesitas 'reportlab'.\n\n"
                "Ejecuta: pip install reportlab",
                "Falta reportlab", parent=self)
            return

        hoy = datetime.now().strftime("%Y-%m-%d")
        arch = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            initialfile=f"historial_precios_{hoy}.pdf",
            filetypes=[("PDF", "*.pdf")])
        if not arch:
            return
        try:
            doc = SimpleDocTemplate(arch, pagesize=letter,
                                    leftMargin=2 * cm, rightMargin=2 * cm,
                                    topMargin=2 * cm, bottomMargin=2 * cm)
            styles = getSampleStyleSheet()
            elements = []
            titulo = Paragraph(
                f"<b>Historial de Precios</b><br/>"
                f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Title"])
            elements.append(titulo)
            elements.append(Spacer(1, 12))

            data = [["Fecha", "Producto", "Código",
                     "Antes", "Ahora", "Dif."]]
            for h in hist:
                dif = h["new_price"] - h["old_price"]
                signo = "+" if dif > 0 else ""
                data.append([
                    h["date"],
                    h.get("product_name", "(eliminado)")[:30],
                    h.get("product_barcode", "") or "-",
                    f"${h['old_price']:,.0f}".replace(",", "."),
                    f"${h['new_price']:,.0f}".replace(",", "."),
                    f"{signo}${dif:,.0f}".replace(",", "."),
                ])
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#305496")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f5f5f5")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.append(table)
            doc.build(elements)
            MD.show_info(f"✅ PDF guardado:\n{arch}", "Listo", parent=self)
        except Exception as e:
            MD.show_error(f"Error al generar PDF: {e}", "Error", parent=self)

    # ============================================================
    # MOVIMIENTO DE CAJA (FASE 6)
    # ============================================================
    def _open_cash_movement(self):
        pop = tk.Toplevel(self)
        pop.title("💵 Movimiento de Caja")
        pop.geometry("420x360")
        pop.transient(self)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="💵 Movimiento de Caja",
                 font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=(15, 4))
        tk.Label(pop, text="Registra retiros, gastos o ingresos manuales.",
                 font=("Arial", 9, "italic"), bg=bg, fg="#a8e6a8").pack(pady=(0, 10))

        tk.Label(pop, text="Tipo:", font=("Arial", 11), bg=bg, fg=fg).pack()
        tipo_var = tk.StringVar(value="retiro")
        tf = tk.Frame(pop, bg=bg)
        tf.pack(pady=5)
        ttk.Radiobutton(tf, text="Retiro (sacar plata)", variable=tipo_var,
                        value="retiro", bootstyle="info").pack(side="left", padx=6)
        ttk.Radiobutton(tf, text="Gasto (pagar algo)", variable=tipo_var,
                        value="gasto", bootstyle="info").pack(side="left", padx=6)
        ttk.Radiobutton(tf, text="Ingreso (plata extra)", variable=tipo_var,
                        value="ingreso", bootstyle="info").pack(side="left", padx=6)

        tk.Label(pop, text="Monto:", font=("Arial", 11), bg=bg, fg=fg).pack(
            pady=(10, 3))
        monto_var = tk.StringVar()
        e = ttk.Entry(pop, textvariable=monto_var, width=20,
                      font=("Arial", 14), justify="center")
        e.pack(pady=5)

        tk.Label(pop, text="Notas (opcional):", font=("Arial", 11),
                 bg=bg, fg=fg).pack(pady=(10, 3))
        notas_var = tk.StringVar()
        ttk.Entry(pop, textvariable=notas_var, width=40).pack(pady=5)

        def guardar():
            try:
                monto = float(monto_var.get().replace("$", "")
                              .replace(".", "").replace(",", ".") or 0)
                if monto <= 0:
                    raise ValueError
            except ValueError:
                MD.show_error("Monto inválido", "Error", parent=pop)
                return
            try:
                self.sale_use_case.register_cash_movement(
                    tipo_var.get(), monto, notas_var.get().strip())
                pop.destroy()
                MD.show_info(
                    f"✅ {tipo_var.get().capitalize()} de "
                    f"${monto:,.0f} registrado.".replace(",", "."),
                    "Listo", parent=self)
            except Exception as ex:
                MD.show_error(f"Error: {ex}", "Error", parent=pop)

        e.bind("<Return>", lambda ev: guardar())
        bf = tk.Frame(pop, bg=bg)
        bf.pack(pady=15)
        ttk.Button(bf, text="💾 Guardar", command=guardar,
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(bf, text="Cancelar",
                   command=pop.destroy).pack(side="left", padx=5)
        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass
        pop.after(80, lambda: e.focus_set() if pop.winfo_exists() else None)

    # ============================================================
    # TAMAÑO DE FUENTE
    # ============================================================
    def _apply_font_size(self):
        size = int(self.db_manager.get_setting("font_size", "11") or 11)
        for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont"):
            try:
                tkfont.nametofont(name).configure(size=size)
            except Exception:
                pass
        self.font_size = size
        self._update_tree_rowheights(size)
        self._setup_dark_green_style()
        try:
            self.style.configure("DarkGreen.TButton",
                                 font=("Arial", size, "bold"))
        except Exception:
            pass

    def _update_tree_rowheights(self, size):
        row_height = size + 14
        try:
            self.style.configure("Treeview", rowheight=row_height,
                                 font=("Arial", size))
            self.style.configure("Treeview.Heading",
                                 font=("Arial", size, "bold"))
        except Exception:
            pass

    def _change_font_size(self, delta):
        new = max(8, min(20, self.font_size + delta))
        self.db_manager.set_setting("font_size", str(new))
        self.font_size = new
        self._apply_font_size()
        try:
            self.payment_view.refresh_cart()
        except Exception:
            pass
        try:
            self.inventory_view.load_products()
        except Exception:
            pass

    # ============================================================
    # UI
    # ============================================================
    def create_widgets(self):
        top_bar = ttk.Frame(self, bootstyle="dark")
        top_bar.pack(fill="x", side="top")

        menubar_wrap = ttk.Frame(top_bar, bootstyle="dark")
        menubar_wrap.pack(side="left", fill="y")
        self.menubar = DarkMenuBar(menubar_wrap,
                                    lambda: self.current_theme == 'darkly')
        self.menubar.pack(side="left")

        try:
            bg = self.style.colors.bg
        except Exception:
            bg = "#1a1a1a"
        right_wrap = tk.Frame(top_bar, bg=bg)
        right_wrap.pack(side="right", padx=(0, 15))
        self.business_header = make_inline_business_header(
            right_wrap, self.db_manager, bg=bg)
        self.business_header["frame"].pack(side="right", pady=6)

        # ==================== MENÚ OPCIONES ====================
        def build_opciones(menu):
            menu.add_command(label="⚙️ Ajustes...",
                             command=self.open_settings_view)
            menu.add_separator()
            menu.add_command(label="🌓 Cambiar Tema", command=self.toggle_theme)
            menu.add_command(label="🏪 Datos del negocio",
                             command=self._open_business_popup)
            menu.add_separator()

            menu.add_command(label="📈 Margen de ganancia por defecto",
                             command=self._configure_default_margin)
            menu.add_command(label="⚠️ Alerta de stock bajo",
                             command=self._configure_low_stock)
            menu.add_command(label="📅 Alertas de vencimiento",
                             command=self._configure_expiry_alerts)
            menu.add_command(label="💾 Respaldo automático",
                             command=self._configure_backup)
            menu.add_separator()

            menu.add_command(label="📜 Exportar Historial de Precios",
                             command=self.export_price_history)
            menu.add_command(label="📦 Exportar Inventario a Excel",
                             command=self.export_inventory_excel)
            menu.add_separator()

            sub_font = tk.Menu(menu, tearoff=0,
                               bg=self.style.colors.bg, fg=self.style.colors.fg,
                               activebackground=self.style.colors.selectbg,
                               activeforeground=self.style.colors.selectfg,
                               font=get_menu_font())
            sub_font.add_command(label="➕ Aumentar letra",
                                 command=lambda: self._change_font_size(1))
            sub_font.add_command(label="➖ Reducir letra",
                                 command=lambda: self._change_font_size(-1))
            menu.add_cascade(label="🔤 Tamaño de letra", menu=sub_font)

            sub_pwd = tk.Menu(menu, tearoff=0,
                              bg=self.style.colors.bg, fg=self.style.colors.fg,
                              activebackground=self.style.colors.selectbg,
                              activeforeground=self.style.colors.selectfg,
                              font=get_menu_font())
            sub_pwd.add_command(label="🆕 Activar/Crear contraseña",
                                command=self._setup_password_first_time)
            sub_pwd.add_command(label="🔐 Cambiar contraseña (admin)",
                                command=self._change_password)
            menu.add_cascade(label="🔑 Contraseña de inicio", menu=sub_pwd)
            menu.add_separator()

            self.tray_var = tk.BooleanVar(
                value=self.db_manager.get_setting("close_to_tray", "0") == "1")
            menu.add_checkbutton(
                label="🔽 Al presionar X ocultar en la bandeja",
                variable=self.tray_var,
                command=self._toggle_close_to_tray)

            self.autostart_var = tk.BooleanVar(value=self._is_autostart_enabled())
            menu.add_checkbutton(label="🚀 Iniciar con Windows",
                                 variable=self.autostart_var,
                                 command=self.toggle_autostart)
            menu.add_separator()

            self.auto_reset_var = tk.BooleanVar(
                value=self.sale_use_case.is_auto_reset_enabled())
            menu.add_checkbutton(
                label="🕛 Reinicio automático a medianoche",
                variable=self.auto_reset_var,
                command=self._toggle_auto_reset)

            menu.add_separator()
            menu.add_command(label="🔄 Reiniciar contador de ventas (#) [manual]",
                             command=self.reset_sales_counter)
            menu.add_separator()

            menu.add_command(label="📄 Reporte del Día (TXT)",
                             command=self.generate_daily_report)
            menu.add_command(label="📊 Reporte del Mes (Excel)",
                             command=self.generate_monthly_report)
            menu.add_separator()
            menu.add_command(label="📤 Exportar Base de Datos",
                             command=self.export_db)
            menu.add_command(label="📥 Importar Base de Datos",
                             command=self.import_db)

        # ==================== MENÚ VENTAS ====================
        def build_ventas(menu):
            menu.add_command(label="📊 Resumen de Ventas (agrupado)",
                             command=self.show_sales_summary)
            menu.add_command(label="💳 Fiados (agrupado por cliente)",
                             command=self.show_credit_sales)
            menu.add_separator()
            menu.add_command(label="↩️ Devoluciones / Notas crédito",
                             command=self.open_returns_window)
            menu.add_separator()
            menu.add_command(label="💰 Cierre de Caja del día",
                             command=self.show_cash_closing)
            menu.add_command(label="💵 Movimiento de Caja (retiro/gasto)",
                             command=self._open_cash_movement)
            menu.add_separator()
            menu.add_command(label="📊 Gráficos y estadísticas",
                             command=self.open_charts)

        # ==================== MENÚ INVENTARIO ====================
        def build_inventario(menu):
            menu.add_command(label="📦 Ver inventario",
                             command=lambda: self.show_page("inventario"))
            menu.add_separator()
            menu.add_command(label="🖨️ Generar etiquetas PDF",
                             command=self._open_labels_from_menu)
            menu.add_command(label="🔔 Alertas de vencimiento",
                             command=self._open_expiry_alerts_from_menu)
            menu.add_command(label="💰 Historial de precios global",
                             command=self._open_price_history_from_menu)

        self.menubar.add_menu("Opciones", build_opciones)
        self.menubar.add_menu("Ventas", build_ventas)
        self.menubar.add_menu("Inventario", build_inventario)

        tabs = ttk.Frame(self, bootstyle="dark")
        tabs.pack(fill="x", padx=10, pady=(6, 0))

        self.btn_pagos = ttk.Button(tabs, text="🛒 PAGOS",
                                    style="DarkGreen.TButton",
                                    command=lambda: self.show_page("pagos"),
                                    width=14)
        self.btn_pagos.pack(side="left", padx=3, pady=2)

        self.btn_inv = ttk.Button(tabs, text="📦 INVENTARIO",
                                  style="DarkGreen.TButton",
                                  command=lambda: self.show_page("inventario"),
                                  width=14)
        self.btn_inv.pack(side="left", padx=3, pady=2)

        self.bind('<Control-Key-1>', lambda e: self.show_page("pagos"))
        self.bind('<Control-Key-2>', lambda e: self.show_page("inventario"))

        self.container = ttk.Frame(self, bootstyle="dark")
        self.container.pack(fill="both", expand=True, padx=10, pady=6)

        self.payment_view = PaymentView(
            self.container, self.product_use_case, self.sale_use_case,
            self.db_manager, lambda: self.current_theme,
            on_business_click=None)
        self.inventory_view = InventoryView(
            self.container, self.product_use_case,
            self.db_manager, lambda: self.current_theme,
            on_business_click=None)
        # Compartir sale_case con inventory_view (para devoluciones y gráficos)
        self.inventory_view._sale_case = self.sale_use_case

        self.current_page = None

    # ============================================================
    # ACCIONES DE MENÚ NUEVAS
    # ============================================================
    def open_settings_view(self):
        try:
            from presentation.views.settings_view import SettingsView
            SettingsView(self, self.db_manager, self.product_use_case,
                         self.sale_use_case,
                         on_theme_change=self._on_theme_change,
                         on_font_change=self._on_font_change,
                         on_business_change=self._on_business_change,
                         on_change_password=self._change_password)
        except Exception as e:
            MD.show_error(f"Error al abrir Ajustes:\n{e}", "Error", parent=self)

    def _on_theme_change(self, theme):
        self.current_theme = theme
        try:
            self.style.theme_use(theme)
        except Exception:
            pass
        self.configure(bg=self.style.colors.bg)
        self._setup_dark_green_style()
        apply_titlebar_theme(self, theme == 'darkly')

    def _on_font_change(self, size):
        try:
            self._apply_font_size()
        except Exception:
            pass

    def _on_business_change(self):
        self._update_window_title()
        self._refresh_business_header()

    def _open_labels_from_menu(self):
        try:
            self.inventory_view.print_labels()
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def _open_expiry_alerts_from_menu(self):
        try:
            self.inventory_view.open_expiry_alerts()
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def _open_price_history_from_menu(self):
        try:
            self.inventory_view.open_price_history_global()
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def open_returns_window(self):
        try:
            from presentation.views.widgets import ReturnsWindow
            ReturnsWindow(self, self.sale_use_case,
                          on_done=self._refresh_all_views)
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def open_charts(self):
        try:
            from presentation.views.widgets import ChartsWindow
            ChartsWindow(self, self.sale_use_case)
        except Exception as e:
            MD.show_error(f"Error: {e}", "Error", parent=self)

    def _refresh_all_views(self):
        try:
            self.inventory_view.load_products()
        except Exception:
            pass
        try:
            self.payment_view._load_recent_sales()
        except Exception:
            pass

    # ============================================================
    # TOGGLE AUTO-RESET
    # ============================================================
    def _toggle_auto_reset(self):
        val = self.auto_reset_var.get()
        self.sale_use_case.enable_auto_reset(val)
        if val:
            MD.show_info("✅ Reinicio automático ACTIVADO.",
                         "Reinicio automático", parent=self)
        else:
            MD.show_info("❌ Reinicio automático DESACTIVADO.",
                         "Reinicio automático", parent=self)

    def _toggle_close_to_tray(self):
        val = "1" if self.tray_var.get() else "0"
        self.db_manager.set_setting("close_to_tray", val)
        if val == "1":
            MD.show_info(
                "✅ Al presionar X la ventana se ocultará en la bandeja del sistema.\n"
                "Para abrirla, haz clic en el ícono (junto al reloj) y elige "
                "'Mostrar MiniPOS'.",
                "Modo bandeja", parent=self)
        else:
            MD.show_info("ℹ️ Al presionar X se pedirá confirmación para salir.",
                         "Modo normal", parent=self)

    # ============================================================
    # RESET CONTADOR CON CONTRASEÑA
    # ============================================================
    def reset_sales_counter(self):
        r1 = MD.yesno(
            "🔄 ¿Reiniciar el contador de ventas?\n\n"
            "Las ventas y fiados NO se borran.\n"
            "Solo el número que aparece como 'Venta #XX' se reiniciará.\n"
            "La próxima venta aparecerá como #01.\n\n"
            "¿Deseas continuar?",
            "Confirmar reinicio", parent=self, default_yes=True)
        if r1 != "Yes":
            return
        pw = self._ask_password_1234(self)
        if not pw:
            return
        try:
            self.sale_use_case.reset_sale_number_counter()
            self.sale_use_case.set_last_reset_date(
                datetime.now().strftime("%Y-%m-%d"))
            MD.show_info(
                "✅ Contador reiniciado.\nLa próxima venta será #01.",
                "Listo", parent=self)
        except Exception as e:
            MD.show_error(f"Error al reiniciar: {e}", "Error", parent=self)

    # ============================================================
    # NAVEGACIÓN
    # ============================================================
    def show_page(self, page):
        for w in (self.payment_view, self.inventory_view):
            w.pack_forget()
        if page == "pagos":
            self.payment_view.pack(fill="both", expand=True)
            self.btn_pagos.configure(style="DarkGreen.TButton")
            self.btn_inv.configure(bootstyle="secondary")
            self.current_page = "pagos"
        else:
            self.inventory_view.load_products()
            self.inventory_view.pack(fill="both", expand=True)
            self.btn_pagos.configure(bootstyle="secondary")
            self.btn_inv.configure(style="DarkGreen.TButton")
            self.current_page = "inventario"

    def toggle_fullscreen(self):
        self.is_fullscreen = not self.is_fullscreen
        self.attributes('-fullscreen', self.is_fullscreen)

    def toggle_theme(self):
        self.current_theme = 'flatly' if self.current_theme == 'darkly' else 'darkly'
        self.style.theme_use(self.current_theme)
        self.configure(bg=self.style.colors.bg)
        self._setup_dark_green_style()
        apply_titlebar_theme(self, self.current_theme == 'darkly')
        self._apply_font_size()
        try:
            bg = self.style.colors.bg
            if hasattr(self, "business_header") and self.business_header:
                frm = self.business_header["frame"]
                frm.configure(bg=bg)
                for child in frm.winfo_children():
                    try:
                        child.configure(bg=bg)
                    except Exception:
                        pass
        except Exception:
            pass

    # ============================================================
    # CIERRE DE CAJA (FASE 6)
    # ============================================================
    def show_cash_closing(self):
        win = tk.Toplevel(self)
        win.title("💰 Cierre de Caja")
        win.geometry("760x820")
        win.transient(self)
        win.configure(bg=self.style.colors.bg)
        win.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        fecha_var = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        resumen_lbl = tk.Label(win, text="", font=("Arial", 11), bg=bg, fg=fg,
                               justify="left", anchor="w", wraplength=680)

        def generar_reporte():
            fecha = fecha_var.get().strip()
            try:
                datetime.strptime(fecha, "%Y-%m-%d")
            except Exception:
                MD.show_error("Fecha inválida. Usa formato YYYY-MM-DD.",
                              "Error", parent=win)
                return
            data = self.sale_use_case.get_cash_closing(fecha)
            lineas = []
            lineas.append(f"📅 CIERRE DE CAJA — {data['date']}")
            lineas.append(f"🕛 Generado: "
                          f"{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
            lineas.append("")
            lineas.append("─── RESUMEN DE VENTAS ───")
            lineas.append(f"🧾 Ventas del día:           "
                          f"{data['cantidad_ventas']}")
            lineas.append(f"📦 Productos vendidos:        "
                          f"{data['productos_vendidos']:g}")
            lineas.append(f"💵 Subtotal bruto:            "
                          f"${data['total_subtotal']:,.0f}".replace(",", "."))
            lineas.append(f"🎁 Descuentos aplicados:     "
                          f"-${data['total_descuento']:,.0f}".replace(",", "."))
            lineas.append(f"💰 TOTAL NETO DEL DÍA:        "
                          f"${data['total_ventas']:,.0f}".replace(",", "."))
            lineas.append("")
            lineas.append("─── DESGLOSE POR MÉTODO ───")
            for metodo in ("Efectivo", "Transferencia", "Tarjeta", "Otro", "Fiado"):
                info = data["metodos"].get(metodo)
                if info:
                    lineas.append(f"  • {metodo:<14} {info['count']:>3} ventas   "
                                  f"${info['total']:>12,.0f}".replace(",", "."))
            otros = [k for k in data["metodos"] if k not in
                     ("Efectivo", "Transferencia", "Tarjeta", "Otro", "Fiado")]
            for k in otros:
                info = data["metodos"][k]
                lineas.append(f"  • {k:<14} {info['count']:>3} ventas   "
                              f"${info['total']:>12,.0f}".replace(",", "."))
            lineas.append("")
            lineas.append("─── MOVIMIENTOS DE CAJA ───")
            lineas.append(f"💵 Efectivo inicial:          "
                          f"${data.get('initial_cash', 0):,.0f}".replace(",", "."))
            lineas.append(f"💰 Ventas en efectivo:        "
                          f"${data['metodos'].get('Efectivo', {}).get('total', 0):,.0f}".replace(",", "."))
            lineas.append(f"💵 Abonos en efectivo:        "
                          f"${data['total_abonos']:,.0f}".replace(",", "."))
            lineas.append(f"➕ Ingresos manuales:         "
                          f"${data.get('ingresos', 0):,.0f}".replace(",", "."))
            lineas.append(f"➖ Retiros:                  "
                          f"-${data.get('retiros', 0):,.0f}".replace(",", "."))
            lineas.append(f"➖ Gastos:                   "
                          f"-${data.get('gastos', 0):,.0f}".replace(",", "."))
            lineas.append(f"↩️ Devoluciones:             "
                          f"-${data.get('total_devuelto', 0):,.0f}".replace(",", "."))
            lineas.append("")
            lineas.append("════════════════════════════════")
            lineas.append(f"💰 EFECTIVO ESPERADO EN CAJA: "
                          f"${data['efectivo_esperado']:,.0f}".replace(",", "."))
            lineas.append("════════════════════════════════")
            if data.get("counted_cash", 0) > 0:
                diff = data.get("difference", 0)
                signo = "+" if diff > 0 else ""
                estado = "SOBRANTE" if diff > 0 else (
                    "FALTANTE" if diff < 0 else "CUADRA")
                lineas.append(f"💵 Efectivo contado:          "
                              f"${data['counted_cash']:,.0f}".replace(",", "."))
                lineas.append(f"📊 Diferencia:               "
                              f"{signo}${diff:,.0f}  ({estado})".replace(",", "."))
            else:
                lineas.append("💵 Efectivo contado:          (sin registrar)")
                lineas.append("👉 Usa el botón 'Contar efectivo' para cuadrar")
            resumen_lbl.configure(text="\n".join(lineas))

        def contar_efectivo():
            fecha = fecha_var.get().strip()
            try:
                datetime.strptime(fecha, "%Y-%m-%d")
            except Exception:
                MD.show_error("Fecha inválida.", "Error", parent=win)
                return
            session = self.sale_use_case.get_or_create_today_cash_session()
            summary = self.sale_use_case.get_cash_session_summary(
                session_id=session["session_id"])
            expected = summary["expected_cash"] if summary else 0.0

            pop = tk.Toplevel(win)
            pop.title("💵 Contar efectivo")
            pop.geometry("440x360")
            pop.transient(win)
            pop.configure(bg=bg)
            pop.withdraw()

            tk.Label(pop, text="💵 CONTAR EFECTIVO",
                     font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=(15, 6))
            tk.Label(pop,
                     text=f"Efectivo esperado: ${expected:,.0f}".replace(",", "."),
                     font=("Arial", 13, "bold"), bg="#0a4d1f", fg="#a8e6a8",
                     padx=12, pady=8).pack(pady=8)

            tk.Label(pop, text="¿Cuánto efectivo hay en caja?",
                     font=("Arial", 11), bg=bg, fg=fg).pack(pady=(10, 4))
            contado_var = tk.StringVar()
            e = ttk.Entry(pop, textvariable=contado_var, width=20,
                          font=("Arial", 16), justify="center")
            e.pack(pady=5)
            e.focus_set()

            diff_lbl = tk.Label(pop, text="", font=("Arial", 12, "bold"),
                                bg=bg, fg=fg)
            diff_lbl.pack(pady=8)

            def actualizar_diff(*args):
                try:
                    valor = float(contado_var.get().replace("$", "")
                                  .replace(".", "").replace(",", ".") or 0)
                except ValueError:
                    valor = 0.0
                diff = valor - expected
                if diff > 0:
                    diff_lbl.configure(
                        text=f"📈 SOBRANTE: ${diff:,.0f}".replace(",", "."),
                        fg="#7dd87d")
                elif diff < 0:
                    diff_lbl.configure(
                        text=f"📉 FALTANTE: ${abs(diff):,.0f}".replace(",", "."),
                        fg="#ff8888")
                else:
                    diff_lbl.configure(text="✅ CUADRA EXACTO", fg="#7dd87d")

            contado_var.trace_add("write", actualizar_diff)

            tk.Label(pop, text="Notas (opcional):",
                     font=("Arial", 10), bg=bg, fg=fg).pack(pady=(8, 3))
            notas_var = tk.StringVar()
            ttk.Entry(pop, textvariable=notas_var, width=40).pack(pady=4)

            def guardar():
                try:
                    valor = float(contado_var.get().replace("$", "")
                                  .replace(".", "").replace(",", ".") or 0)
                except ValueError:
                    MD.show_error("Valor inválido", "Error", parent=pop)
                    return
                try:
                    self.sale_use_case.update_counted_cash(
                        session["session_id"], valor, notas_var.get().strip())
                    pop.destroy()
                    generar_reporte()
                    MD.show_info("✅ Efectivo contado registrado.", "Listo",
                                 parent=win)
                except Exception as ex:
                    MD.show_error(f"Error: {ex}", "Error", parent=pop)

            bf = tk.Frame(pop, bg=bg)
            bf.pack(pady=15)
            ttk.Button(bf, text="💾 Guardar", command=guardar,
                       style="DarkGreen.TButton").pack(side="left", padx=5)
            ttk.Button(bf, text="Cancelar",
                       command=pop.destroy).pack(side="left", padx=5)
            show_popup_smooth(pop)
            try:
                pop.grab_set()
                pop.focus_force()
            except Exception:
                pass

        def exportar_txt():
            fecha = fecha_var.get().strip()
            try:
                datetime.strptime(fecha, "%Y-%m-%d")
            except Exception:
                MD.show_error("Fecha inválida.", "Error", parent=win)
                return
            arch = filedialog.asksaveasfilename(
                defaultextension=".txt",
                initialfile=f"cierre_caja_{fecha}.txt",
                filetypes=[("Texto", "*.txt")],
                parent=win)
            if not arch:
                return
            try:
                with open(arch, "w", encoding="utf-8") as f:
                    f.write(resumen_lbl.cget("text"))
                MD.show_info(f"Guardado:\n{arch}", "Listo", parent=win)
            except Exception as e:
                MD.show_error(f"Error: {e}", "Error", parent=win)

        container = tk.Frame(win, bg=bg)
        container.pack(fill="both", expand=True, padx=10, pady=10)

        def build_content(parent):
            tk.Label(parent, text="💰 CIERRE DE CAJA",
                     font=("Arial", 18, "bold"), bg=bg, fg=fg).pack(pady=(10, 6))
            row = tk.Frame(parent, bg=bg)
            row.pack(pady=6)
            tk.Label(row, text="Fecha (YYYY-MM-DD):",
                     font=("Arial", 11), bg=bg, fg=fg).pack(side="left", padx=5)
            ttk.Entry(row, textvariable=fecha_var, width=15,
                      font=("Arial", 12), justify="center").pack(side="left", padx=5)
            ttk.Button(row, text="🔍 Generar",
                       command=generar_reporte,
                       style="DarkGreen.TButton").pack(side="left", padx=8)
            resumen_lbl.pack(fill="both", expand=True, padx=10, pady=10)

        def build_bottom(parent):
            ttk.Button(parent, text="💵 Contar efectivo",
                       command=contar_efectivo,
                       style="DarkGreen.TButton").pack(side="left", padx=6)
            ttk.Button(parent, text="📄 Exportar a TXT",
                       command=exportar_txt,
                       bootstyle="info").pack(side="left", padx=6)
            ttk.Button(parent, text="Cerrar",
                       command=win.destroy).pack(side="right", padx=6)

        make_scrollable(container, build_content, build_bottom, bg=bg)
        generar_reporte()
        show_popup_smooth(win)
        try:
            win.grab_set()
            win.focus_force()
        except Exception:
            pass

    # ============================================================
    # IMPRIMIR TICKET DE UNA VENTA
    # ============================================================
    def _print_ticket_for_sale(self, sale_id, parent=None):
        try:
            archivo = self.sale_use_case.generate_ticket_pdf(
                sale_id, base_dir=self.base_dir)
            if not archivo:
                MD.show_error("No se pudo generar el ticket.\n\n"
                              "Verifica que reportlab esté instalado.",
                              "Error", parent=parent or self)
                return
            r = MD.yesno(
                f"✅ Ticket generado:\n\n{archivo}\n\n¿Quieres abrirlo ahora?",
                "Ticket listo", parent=parent or self)
            if r == "Yes":
                try:
                    if sys.platform.startswith("win"):
                        os.startfile(archivo)
                    elif sys.platform == "darwin":
                        os.system(f"open \"{archivo}\"")
                    else:
                        os.system(f"xdg-open \"{archivo}\"")
                except Exception:
                    pass
        except Exception as e:
            MD.show_error(f"Error al generar ticket: {e}", "Error",
                          parent=parent or self)

    # ============================================================
    # PEDIR CONTRASEÑA 1234
    # ============================================================
    def _ask_password_1234(self, parent):
        pop = tk.Toplevel(parent)
        pop.title("Contraseña requerida")
        pop.geometry("340x220")
        pop.transient(parent)
        pop.configure(bg=self.style.colors.bg)
        pop.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        tk.Label(pop, text="🔒 Contraseña de administrador:",
                 font=("Arial", 12, "bold"), bg=bg, fg=fg).pack(pady=15)
        v = tk.StringVar()
        e = ttk.Entry(pop, textvariable=v, show="•", width=20,
                      font=("Arial", 14), justify="center")
        e.pack(pady=5)
        ok = {"v": False}

        def ver(ev=None):
            if v.get() == "1234":
                ok["v"] = True
                pop.destroy()
            else:
                MD.show_error("Contraseña incorrecta", "Error", parent=pop)
                v.set("")
                try:
                    e.focus_set()
                    e.select_range(0, tk.END)
                except Exception:
                    pass
            return "break"

        def cancelar(ev=None):
            pop.destroy()
            return "break"

        e.bind("<Return>", ver)
        e.bind("<KP_Enter>", ver)
        pop.bind("<Escape>", cancelar)

        ttk.Button(pop, text="Aceptar", command=ver,
                   style="DarkGreen.TButton").pack(pady=15)

        show_popup_smooth(pop)
        try:
            pop.grab_set()
            pop.focus_force()
        except Exception:
            pass

        def set_focus():
            try:
                if pop.winfo_exists():
                    e.focus_set()
                    e.select_range(0, tk.END)
            except Exception:
                pass
        pop.after(50, set_focus)
        pop.after(250, set_focus)

        parent.wait_window(pop)
        return ok["v"]

    # ============================================================
    # RESUMEN DE VENTAS (COMPLETO)
    # ============================================================
    def show_sales_summary(self):
        win = tk.Toplevel(self)
        win.title("Resumen de Ventas")
        win.geometry("1200x780")
        win.transient(self)
        win.configure(bg=self.style.colors.bg)
        win.withdraw()
        bg = self.style.colors.bg
        fg = self.style.colors.fg

        container = tk.Frame(win, bg=bg)
        container.pack(fill="both", expand=True)

        grupos_map = {}
        marcados = set()
        refs = {}

        def build_content(parent):
            tk.Label(parent, text="📊 RESUMEN DE VENTAS (agrupado por cliente)",
                     font=("Arial", 18, "bold"), bg=bg, fg=fg).pack(pady=(12, 6))

            # ---- CARDS dinámicas ----
            cards = tk.Frame(parent, bg=bg)
            cards.pack(pady=5)
            cards_labels = {}
            refs["cards"] = cards
            refs["cards_labels"] = cards_labels

            s = self.sale_use_case.get_summary()
            datos = [
                ("HOY", s["hoy"], "#0d6efd"),
                ("MES", s["mes"], "#0d6efd"),
                ("TOTAL", s["total"], "#0a4d1f"),
                ("FIADOS", s["fiados"], "#d97706"),
            ]
            for i, (titulo, (cnt, tot), color) in enumerate(datos):
                c = tk.Frame(cards, bg=color, padx=18, pady=10)
                c.grid(row=0, column=i, padx=8)
                tk.Label(c, text=titulo, font=("Arial", 12, "bold"),
                         bg=color, fg="#ffffff").pack()
                lc = tk.Label(c, text=f"{cnt} ventas", font=("Arial", 11),
                              bg=color, fg="#ffffff")
                lc.pack()
                lt = tk.Label(c, text=f"${tot:,.0f}".replace(",", "."),
                              font=("Arial", 16, "bold"),
                              bg=color, fg="#ffffff")
                lt.pack()
                cards_labels[titulo] = (lc, lt)

            # ---- FILTROS ----
            filt_frame = tk.Frame(parent, bg=bg)
            filt_frame.pack(pady=5)
            filtro_var = tk.StringVar(value="all")
            refs["filtro_var"] = filtro_var
            for val, txt in [("all", "Todas"), ("today", "Hoy"),
                             ("month", "Este mes"), ("range", "Rango")]:
                ttk.Radiobutton(filt_frame, text=txt, variable=filtro_var,
                                value=val, bootstyle="info",
                                command=lambda: on_filter_change()).pack(
                                    side="left", padx=8)

            rango_frame = tk.Frame(parent, bg=bg)
            rango_frame.pack(pady=3)
            tk.Label(rango_frame, text="Desde:", font=("Arial", 10),
                     bg=bg, fg=fg).pack(side="left", padx=(0, 3))
            hoy = datetime.now().strftime("%Y-%m-%d")
            hace7 = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d")
            desde_var = tk.StringVar(value=hace7)
            ttk.Entry(rango_frame, textvariable=desde_var, width=12,
                      font=("Arial", 10), justify="center").pack(side="left", padx=3)
            tk.Label(rango_frame, text="Hasta:", font=("Arial", 10),
                     bg=bg, fg=fg).pack(side="left", padx=(10, 3))
            hasta_var = tk.StringVar(value=hoy)
            ttk.Entry(rango_frame, textvariable=hasta_var, width=12,
                      font=("Arial", 10), justify="center").pack(side="left", padx=3)
            ttk.Button(rango_frame, text="Aplicar",
                       command=lambda: on_filter_change(),
                       bootstyle="info").pack(side="left", padx=8)
            refs["desde_var"] = desde_var
            refs["hasta_var"] = hasta_var
            refs["rango_frame"] = rango_frame

            # ---- TREEVIEW ----
            tree_frame = ttk.Frame(parent, bootstyle="dark")
            tree_frame.pack(fill="both", expand=True, padx=15, pady=8)

            tree_frame2, tree = make_scrolled_treeview(
                tree_frame,
                columns=("Sel", "Cliente", "Ventas", "Total", "Pagado", "Pendiente"),
                headings=[
                    ("Sel", "☐", 40, "center"),
                    ("Cliente", "Cliente", 240, "center"),
                    ("Ventas", "# Ventas", 90, "center"),
                    ("Total", "Total", 140, "center"),
                    ("Pagado", "Pagado", 140, "center"),
                    ("Pendiente", "Pendiente", 140, "center"),
                ],
                bootstyle="dark")
            tree_frame2.pack(fill="both", expand=True)

            tree.heading("Sel", text="☐", command=toggle_all)
            refs["tree"] = tree

            tree.bind("<Double-1>", ver_detalle)
            tree.bind("<Button-3>", on_right_click)
            tree.bind("<Button-1>", on_click, add="+")
            win.bind("<Control-F12>", on_ctrl_f12)
            win.bind("<Shift-F12>", on_shift_f12)
            tree.bind("<Control-F12>", on_ctrl_f12)
            tree.bind("<Shift-F12>", on_shift_f12)

        def build_bottom(parent):
            ttk.Button(parent, text="📋 Más detalles",
                       command=lambda: ver_detalle(),
                       bootstyle="info").pack(side="left", padx=5)
            ttk.Button(parent, text="☑ Marcar todos",
                       command=lambda: toggle_all(),
                       bootstyle="secondary").pack(side="left", padx=5)
            ttk.Button(parent, text="🗑️ Eliminar marcados",
                       command=lambda: eliminar_marcados(),
                       bootstyle="danger").pack(side="left", padx=5)
            ttk.Button(parent, text="🔄 Refrescar",
                       command=lambda: recargar(),
                       bootstyle="secondary").pack(side="left", padx=5)

        def actualizar_cards():
            try:
                s = self.sale_use_case.get_summary()
                datos = {"HOY": s["hoy"], "MES": s["mes"],
                         "TOTAL": s["total"], "FIADOS": s["fiados"]}
                labels = refs.get("cards_labels", {})
                for titulo, (cnt, tot) in datos.items():
                    if titulo in labels:
                        lc, lt = labels[titulo]
                        lc.configure(text=f"{cnt} ventas")
                        lt.configure(text=f"${tot:,.0f}".replace(",", "."))
            except Exception:
                pass

        def actualizar_heading_sel():
            try:
                tree = refs.get("tree")
                if not tree:
                    return
                if not grupos_map or len(marcados) == 0:
                    texto = "☐"
                elif len(marcados) >= len(grupos_map):
                    texto = "☑"
                else:
                    texto = "◪"
                tree.heading("Sel", text=texto)
            except Exception:
                pass

        def toggle_mark(iid):
            if not iid or iid not in grupos_map:
                return
            tree = refs.get("tree")
            if iid in marcados:
                marcados.discard(iid)
                try:
                    tree.set(iid, "Sel", "☐")
                except Exception:
                    pass
            else:
                marcados.add(iid)
                try:
                    tree.set(iid, "Sel", "☑")
                except Exception:
                    pass
            actualizar_heading_sel()

        def toggle_all():
            tree = refs.get("tree")
            if not tree or not grupos_map:
                return
            if len(marcados) >= len(grupos_map):
                marcados.clear()
                for iid in grupos_map:
                    try:
                        tree.set(iid, "Sel", "☐")
                    except Exception:
                        pass
            else:
                marcados.clear()
                for iid in grupos_map:
                    marcados.add(iid)
                    try:
                        tree.set(iid, "Sel", "☑")
                    except Exception:
                        pass
            actualizar_heading_sel()

        def desmarcar_todos():
            tree = refs.get("tree")
            marcados.clear()
            if tree:
                for iid in grupos_map:
                    try:
                        tree.set(iid, "Sel", "☐")
                    except Exception:
                        pass
            actualizar_heading_sel()

        def restaurar_foco():
            try:
                if win.winfo_exists():
                    win.lift()
                    win.focus_force()
            except Exception:
                pass

        def on_filter_change():
            try:
                filtro_var = refs.get("filtro_var")
                rango_frame = refs.get("rango_frame")
                if filtro_var and rango_frame:
                    if filtro_var.get() == "range":
                        rango_frame.pack_forget()
                        rango_frame.pack(pady=3)
                    else:
                        rango_frame.pack_forget()
            except Exception:
                pass
            recargar()

        def recargar():
            tree = refs.get("tree")
            if not tree:
                return
            for r in tree.get_children():
                tree.delete(r)
            grupos_map.clear()
            marcados.clear()

            filtro_var = refs.get("filtro_var")
            filtro = filtro_var.get() if filtro_var else "all"

            if filtro == "today":
                sales = self.sale_use_case.get_sales_by_day(
                    datetime.now().strftime("%Y-%m-%d"))
            elif filtro == "month":
                sales = self.sale_use_case.get_sales_by_month(
                    datetime.now().strftime("%Y-%m"))
            elif filtro == "range":
                desde = refs.get("desde_var").get().strip() if refs.get("desde_var") else ""
                hasta = refs.get("hasta_var").get().strip() if refs.get("hasta_var") else ""
                try:
                    datetime.strptime(desde, "%Y-%m-%d")
                    datetime.strptime(hasta, "%Y-%m-%d")
                except Exception:
                    MD.show_error("Fechas inválidas. Usa YYYY-MM-DD.",
                                  "Error", parent=win)
                    sales = []
                else:
                    sales = self.sale_use_case.get_sales_by_range(desde, hasta)
            else:
                sales = self.sale_use_case.get_all_sales(limit=1000)

            grupos = self.sale_use_case.group_sales_by_customer(sales)
            for g in grupos:
                iid = tree.insert("", "end", values=(
                    "☐", g["customer_name"], g["count"],
                    f"${g['total']:,.0f}".replace(",", "."),
                    f"${g['paid']:,.0f}".replace(",", "."),
                    f"${g['pending']:,.0f}".replace(",", ".")))
                grupos_map[iid] = g
            actualizar_heading_sel()
            actualizar_cards()

        def ver_detalle(event=None):
            tree = refs.get("tree")
            if not tree:
                return
            sel = tree.selection()
            if not sel:
                MD.show_warning("Selecciona un cliente primero.",
                                "Sin selección", parent=win)
                restaurar_foco()
                return
            g = grupos_map.get(sel[0])
            if not g:
                return
            det = tk.Toplevel(win)
            det.title(f"Detalle - {g['customer_name']}")
            det.geometry("1050x580")
            det.transient(win)
            det.configure(bg=bg)
            det.withdraw()

            tk.Label(det, text=f"👤 {g['customer_name']}  |  "
                               f"{g['count']} ventas",
                     font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=8)
            tk.Label(det,
                     text=f"Total: ${g['total']:,.0f}   |   "
                          f"Pagado: ${g['paid']:,.0f}   |   "
                          f"Pendiente: ${g['pending']:,.0f}".replace(",", "."),
                     font=("Arial", 11), bg=bg, fg=fg).pack(pady=4)

            f2 = ttk.Frame(det, bootstyle="dark")
            f2.pack(fill="both", expand=True, padx=12, pady=6)
            tf2, t2 = make_scrolled_treeview(
                f2,
                columns=("ID", "Fecha", "Método", "Productos",
                         "Subtotal", "Desc.", "Total", "Estado"),
                headings=[
                    ("ID", "#", 55, "center"),
                    ("Fecha", "Fecha", 135, "center"),
                    ("Método", "Método", 100, "center"),
                    ("Productos", "Productos", 220, "center"),
                    ("Subtotal", "Subtotal", 90, "center"),
                    ("Desc.", "Desc.", 75, "center"),
                    ("Total", "Total", 90, "center"),
                    ("Estado", "Estado", 110, "center"),
                ],
                bootstyle="dark")
            tf2.pack(fill="both", expand=True)

            for sale in sorted(g["sales"], key=lambda x: x.date):
                resumen_items = ", ".join(
                    f"{it.quantity:g}x {it.product_name[:20]}"
                    for it in sale.items[:3])
                if len(sale.items) > 3:
                    resumen_items += f" (+{len(sale.items) - 3} más)"
                estado = "✅ Pagado"
                if sale.is_credit:
                    estado = "💳 Fiado" if not sale.is_paid else "✅ Fiado pagado"
                sub = getattr(sale, "subtotal", sale.total)
                desc = getattr(sale, "discount", 0.0)
                t2.insert("", "end", values=(
                    f"#{sale.display_number:02d}",
                    sale.date, sale.payment_method, resumen_items,
                    f"${sub:,.0f}".replace(",", "."),
                    f"-${desc:,.0f}".replace(",", ".") if desc > 0 else "-",
                    f"${sale.total:,.0f}".replace(",", "."),
                    estado))

            bf2 = tk.Frame(det, bg=bg)
            bf2.pack(side="bottom", pady=8)

            def imprimir_seleccion():
                sel2 = t2.selection()
                if not sel2:
                    MD.show_warning("Selecciona una venta primero.",
                                    "Sin selección", parent=det)
                    return
                vals = t2.item(sel2[0], 'values')
                try:
                    display_num = int(vals[0].replace("#", ""))
                except Exception:
                    return
                for sale in g["sales"]:
                    if sale.display_number == display_num:
                        self._print_ticket_for_sale(sale.sale_id, det)
                        break

            ttk.Button(bf2, text="🧾 Imprimir ticket",
                       command=imprimir_seleccion,
                       style="DarkGreen.TButton").pack(side="left", padx=5)
            ttk.Button(bf2, text="Cerrar",
                       command=lambda: [det.destroy(), restaurar_foco()]).pack(
                           side="right", padx=5)

            show_popup_smooth(det)
            try:
                det.grab_set()
                det.focus_force()
            except Exception:
                pass

        def on_click(event):
            try:
                tree = refs.get("tree")
                region = tree.identify("region", event.x, event.y)
                if region != "cell":
                    return
                col = tree.identify_column(event.x)
                row = tree.identify_row(event.y)
                if not row:
                    return
                if col == "#1":
                    toggle_mark(row)
            except Exception:
                pass

        def eliminar_marcados():
            if not marcados:
                MD.show_warning("No hay clientes marcados.",
                                "Sin selección", parent=win)
                restaurar_foco()
                return
            total_ventas = sum(grupos_map[i]["count"] for i in marcados
                               if i in grupos_map)
            total_dinero = sum(grupos_map[i]["total"] for i in marcados
                               if i in grupos_map)
            if MD.yesno(
                    f"⚠️ ¿Eliminar TODAS las ventas de los "
                    f"{len(marcados)} clientes marcados?\n\n"
                    f"Se eliminarán {total_ventas} ventas por un total de "
                    f"${total_dinero:,.0f}.\n"
                    f"Esto restaurará el stock de todos los productos.".replace(
                        ",", "."),
                    "Confirmar eliminación total", parent=win) != "Yes":
                restaurar_foco()
                return
            pw = self._ask_password_1234(win)
            if not pw:
                restaurar_foco()
                return
            total_borradas = 0
            for iid in list(marcados):
                g = grupos_map.get(iid)
                if not g:
                    continue
                for s in g["sales"]:
                    self.sale_use_case.delete_sale(s.sale_id)
                    total_borradas += 1
            MD.show_info(
                f"✅ {total_borradas} ventas eliminadas y stock restaurado.",
                "Listo", parent=win)
            recargar()
            if self.current_page == "inventario":
                self.inventory_view.load_products()
            restaurar_foco()

        def on_ctrl_f12(event=None):
            tree = refs.get("tree")
            if not tree:
                return
            iid = None
            if marcados:
                iid = next(iter(marcados))
            else:
                sel = tree.selection()
                if sel:
                    iid = sel[0]
            if not iid:
                MD.show_warning("Marca o selecciona un cliente primero.",
                                "Sin selección", parent=win)
                restaurar_foco()
                return
            g = grupos_map.get(iid)
            if not g or not g["sales"]:
                return
            venta_reciente = max(g["sales"], key=lambda x: x.sale_id)
            sid = venta_reciente.sale_id
            display = venta_reciente.display_number
            if MD.yesno(
                    f"⚠️ ¿Eliminar la venta más reciente de "
                    f"'{g['customer_name']}'?\n\n"
                    f"Venta #{display:02d} - ${venta_reciente.total:,.0f}\n"
                    f"Esto restaurará el stock.".replace(",", "."),
                    "Confirmar eliminación", parent=win) != "Yes":
                restaurar_foco()
                return
            pw = self._ask_password_1234(win)
            if not pw:
                restaurar_foco()
                return
            self.sale_use_case.delete_sale(sid)
            MD.show_info(
                f"Venta #{display:02d} eliminada y stock restaurado.",
                "Listo", parent=win)
            recargar()
            if self.current_page == "inventario":
                self.inventory_view.load_products()
            restaurar_foco()

        def on_shift_f12(event=None):
            eliminar_marcados()

        def on_right_click(event):
            try:
                tree = refs.get("tree")
                row = tree.identify_row(event.y)
                if row:
                    tree.selection_set(row)
                    tree.focus(row)
                style = ttk.Style()
                m = tk.Menu(win, tearoff=0,
                            bg=style.colors.bg, fg=style.colors.fg,
                            activebackground=style.colors.selectbg,
                            activeforeground=style.colors.selectfg,
                            bd=1, relief="solid",
                            font=get_menu_font())
                m.add_command(label="👁️  Ver detalle", command=ver_detalle)
                m.add_separator()
                if row:
                    marca = "☑ Desmarcar" if row in marcados else "☐ Marcar"
                    m.add_command(label=marca,
                                  command=lambda r=row: toggle_mark(r))
                m.add_command(label="☑ Marcar todos", command=toggle_all)
                m.add_command(label="☐ Desmarcar todos",
                              command=desmarcar_todos)
                m.add_separator()
                m.add_command(
                    label=f"🗑️  Eliminar marcados ({len(marcados)})",
                    command=eliminar_marcados)
                try:
                    m.tk_popup(event.x_root, event.y_root)
                finally:
                    m.grab_release()
            except Exception:
                pass

        make_scrollable(container, build_content, build_bottom, bg=bg)
        try:
            refs["rango_frame"].pack_forget()
        except Exception:
            pass
        recargar()

        show_popup_smooth(win)
        try:
            win.grab_set()
            win.focus_force()
        except Exception:
            pass
            # ============================================================
# FIADOS — versión completa con checkboxes y hotkeys
# ============================================================
def show_credit_sales(self):
    win = tk.Toplevel(self)
    win.title("Fiados - Cuentas por cobrar")
    win.geometry("1150x720")
    win.transient(self)
    win.configure(bg=self.style.colors.bg)
    win.withdraw()
    bg = self.style.colors.bg
    fg = self.style.colors.fg

    container = tk.Frame(win, bg=bg)
    container.pack(fill="both", expand=True)

    grupos_map = {}
    marcados = set()
    refs = {}

    def build_content(parent):
        tk.Label(parent,
                 text="💳 CUENTAS POR COBRAR (agrupado por cliente)",
                 font=("Arial", 18, "bold"), bg=bg, fg=fg).pack(pady=(12, 6))

        resumen_lbl = tk.Label(parent, text="",
                                font=("Arial", 12, "bold"),
                                bg=bg, fg=fg)
        resumen_lbl.pack(pady=4)
        refs["resumen_lbl"] = resumen_lbl

        tree_frame = ttk.Frame(parent, bootstyle="dark")
        tree_frame.pack(fill="both", expand=True, padx=15, pady=8)

        tf2, tree = make_scrolled_treeview(
            tree_frame,
            columns=("Sel", "Cliente", "Fiados", "Total",
                     "Abonado", "Pendiente"),
            headings=[
                ("Sel", "☐", 40, "center"),
                ("Cliente", "Cliente", 240, "center"),
                ("Fiados", "# Fiados", 90, "center"),
                ("Total", "Total fiado", 140, "center"),
                ("Abonado", "Abonado", 140, "center"),
                ("Pendiente", "Pendiente", 140, "center"),
            ],
            bootstyle="dark")
        tf2.pack(fill="both", expand=True)

        tree.heading("Sel", text="☐", command=toggle_all)
        refs["tree"] = tree

        tree.bind("<Double-1>", ver_detalle)
        tree.bind("<Button-1>", on_click, add="+")
        tree.bind("<Button-3>", on_right_click)
        win.bind("<Control-F12>", on_ctrl_f12)
        win.bind("<Shift-F12>", on_shift_f12)
        tree.bind("<Control-F12>", on_ctrl_f12)
        tree.bind("<Shift-F12>", on_shift_f12)

    def build_bottom(parent):
        ttk.Button(parent, text="📋 Más detalles",
                   command=lambda: ver_detalle(),
                   bootstyle="info").pack(side="left", padx=5)
        ttk.Button(parent, text="💵 Abonar (más antigua)",
                   command=lambda: abonar(),
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(parent, text="✅ Marcar como pagado",
                   command=lambda: marcar_pagado(),
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(parent, text="☑ Marcar todos",
                   command=lambda: toggle_all(),
                   bootstyle="secondary").pack(side="left", padx=5)
        ttk.Button(parent, text="🗑️ Eliminar marcados",
                   command=lambda: eliminar_marcados(),
                   bootstyle="danger").pack(side="left", padx=5)
        ttk.Button(parent, text="🔄 Refrescar",
                   command=lambda: recargar(),
                   bootstyle="secondary").pack(side="left", padx=5)

    def actualizar_heading_sel():
        try:
            tree = refs.get("tree")
            if not tree:
                return
            if not grupos_map or len(marcados) == 0:
                texto = "☐"
            elif len(marcados) >= len(grupos_map):
                texto = "☑"
            else:
                texto = "◪"
            tree.heading("Sel", text=texto)
        except Exception:
            pass

    def toggle_mark(iid):
        if not iid or iid not in grupos_map:
            return
        tree = refs.get("tree")
        if iid in marcados:
            marcados.discard(iid)
            try:
                tree.set(iid, "Sel", "☐")
            except Exception:
                pass
        else:
            marcados.add(iid)
            try:
                tree.set(iid, "Sel", "☑")
            except Exception:
                pass
        actualizar_heading_sel()

    def toggle_all():
        tree = refs.get("tree")
        if not tree or not grupos_map:
            return
        if len(marcados) >= len(grupos_map):
            marcados.clear()
            for iid in grupos_map:
                try:
                    tree.set(iid, "Sel", "☐")
                except Exception:
                    pass
        else:
            marcados.clear()
            for iid in grupos_map:
                marcados.add(iid)
                try:
                    tree.set(iid, "Sel", "☑")
                except Exception:
                    pass
        actualizar_heading_sel()

    def desmarcar_todos():
        tree = refs.get("tree")
        marcados.clear()
        if tree:
            for iid in grupos_map:
                try:
                    tree.set(iid, "Sel", "☐")
                except Exception:
                    pass
        actualizar_heading_sel()

    def restaurar_foco():
        try:
            if win.winfo_exists():
                win.lift()
                win.focus_force()
        except Exception:
            pass

    def recargar():
        tree = refs.get("tree")
        if not tree:
            return
        for r in tree.get_children():
            tree.delete(r)
        grupos_map.clear()
        marcados.clear()
        ventas = self.sale_use_case.get_credit_sales(only_unpaid=True)
        ventas = [v for v in ventas if v.pending() > 0.01]
        grupos = self.sale_use_case.group_sales_by_customer(ventas)
        total_global = 0.0
        for g in grupos:
            total_global += g["pending"]
            iid = tree.insert("", "end", values=(
                "☐", g["customer_name"], g["count"],
                f"${g['total']:,.0f}".replace(",", "."),
                f"${g['paid']:,.0f}".replace(",", "."),
                f"${g['pending']:,.0f}".replace(",", ".")))
            grupos_map[iid] = g
        try:
            refs["resumen_lbl"].configure(
                text=f"👥 {len(grupos)} clientes con deuda   |   "
                     f"💰 Total por cobrar: ${total_global:,.0f}".replace(
                         ",", "."))
        except Exception:
            pass
        actualizar_heading_sel()

    def ver_detalle(event=None):
        tree = refs.get("tree")
        if not tree:
            return
        sel = tree.selection()
        if not sel:
            MD.show_warning("Selecciona un cliente primero.",
                            "Sin selección", parent=win)
            restaurar_foco()
            return
        g = grupos_map.get(sel[0])
        if not g:
            return
        det = tk.Toplevel(win)
        det.title(f"Fiados de {g['customer_name']}")
        det.geometry("950x560")
        det.transient(win)
        det.configure(bg=bg)
        det.withdraw()

        tk.Label(det, text=f"👤 {g['customer_name']}",
                 font=("Arial", 14, "bold"), bg=bg, fg=fg).pack(pady=8)
        tk.Label(det,
                 text=f"Total: ${g['total']:,.0f}   |   "
                      f"Abonado: ${g['paid']:,.0f}   |   "
                      f"Pendiente: ${g['pending']:,.0f}".replace(",", "."),
                 font=("Arial", 11), bg=bg, fg=fg).pack(pady=4)

        f2 = ttk.Frame(det, bootstyle="dark")
        f2.pack(fill="both", expand=True, padx=12, pady=6)
        tf2, t2 = make_scrolled_treeview(
            f2,
            columns=("ID", "Fecha", "Productos", "Total",
                     "Abonado", "Pendiente"),
            headings=[
                ("ID", "#", 60, "center"),
                ("Fecha", "Fecha", 130, "center"),
                ("Productos", "Productos", 280, "center"),
                ("Total", "Total", 90, "center"),
                ("Abonado", "Abonado", 90, "center"),
                ("Pendiente", "Pendiente", 100, "center"),
            ],
            bootstyle="dark")
        tf2.pack(fill="both", expand=True)

        for sale in sorted(g["sales"], key=lambda x: x.date):
            resumen_items = ", ".join(
                f"{it.quantity:g}x {it.product_name[:18]}"
                for it in sale.items[:3])
            if len(sale.items) > 3:
                resumen_items += f" (+{len(sale.items) - 3} más)"
            t2.insert("", "end", values=(
                f"#{sale.display_number:02d}",
                sale.date, resumen_items,
                f"${sale.total:,.0f}".replace(",", "."),
                f"${sale.amount_paid:,.0f}".replace(",", "."),
                f"${sale.pending():,.0f}".replace(",", ".")))

        bf2 = tk.Frame(det, bg=bg)
        bf2.pack(side="bottom", pady=8)
        ttk.Button(bf2, text="Cerrar",
                   command=lambda: [det.destroy(), restaurar_foco()]).pack()

        show_popup_smooth(det)
        try:
            det.grab_set()
            det.focus_force()
        except Exception:
            pass

    def on_click(event):
        try:
            tree = refs.get("tree")
            region = tree.identify("region", event.x, event.y)
            if region != "cell":
                return
            col = tree.identify_column(event.x)
            row = tree.identify_row(event.y)
            if not row:
                return
            if col == "#1":
                toggle_mark(row)
        except Exception:
            pass

    def abonar():
        tree = refs.get("tree")
        if not tree:
            return
        iid = None
        if marcados:
            iid = next(iter(marcados))
        else:
            sel = tree.selection()
            if sel:
                iid = sel[0]
        if not iid:
            MD.show_warning("Selecciona un cliente primero.",
                            "Sin selección", parent=win)
            restaurar_foco()
            return
        g = grupos_map.get(iid)
        if not g:
            return
        ventas_pend = [v for v in g["sales"] if v.pending() > 0.01]
        if not ventas_pend:
            MD.show_info("Este cliente no tiene deudas pendientes.",
                         "Listo", parent=win)
            restaurar_foco()
            return
        venta = sorted(ventas_pend, key=lambda x: x.sale_id)[0]
        pendiente = venta.pending()
        self._abonar_dialog(win, venta, pendiente, recargar)

    def marcar_pagado():
        tree = refs.get("tree")
        if not tree:
            return
        if not marcados and not tree.selection():
            MD.show_warning("Marca o selecciona un cliente primero.",
                            "Sin selección", parent=win)
            restaurar_foco()
            return
        iids = list(marcados) if marcados else list(tree.selection())
        total_marcar = sum(grupos_map[i]["pending"] for i in iids
                           if i in grupos_map)
        if MD.yesno(
                f"¿Marcar TODOS los fiados pendientes de {len(iids)} "
                f"cliente(s) como PAGADOS?\n"
                f"Total a marcar: ${total_marcar:,.0f}".replace(",", "."),
                "Confirmar", parent=win) != "Yes":
            restaurar_foco()
            return
        for iid in iids:
            g = grupos_map.get(iid)
            if not g:
                continue
            for v in g["sales"]:
                if v.pending() > 0.01:
                    self.sale_use_case.mark_as_paid(v.sale_id)
        recargar()
        MD.show_info("Fiados marcados como pagados.", "Listo", parent=win)
        restaurar_foco()

    def eliminar_marcados():
        if not marcados:
            MD.show_warning("No hay clientes marcados.",
                            "Sin selección", parent=win)
            restaurar_foco()
            return
        total_fiados = sum(grupos_map[i]["count"] for i in marcados
                           if i in grupos_map)
        total_dinero = sum(grupos_map[i]["total"] for i in marcados
                           if i in grupos_map)
        if MD.yesno(
                f"⚠️ ¿Eliminar TODOS los fiados de los {len(marcados)} "
                f"clientes marcados?\n\n"
                f"Se eliminarán {total_fiados} fiados por un total de "
                f"${total_dinero:,.0f}.\n"
                f"Esto restaurará el stock de todos los productos.".replace(
                    ",", "."),
                "Confirmar eliminación total", parent=win) != "Yes":
            restaurar_foco()
            return
        pw = self._ask_password_1234(win)
        if not pw:
            restaurar_foco()
            return
        total_borradas = 0
        for iid in list(marcados):
            g = grupos_map.get(iid)
            if not g:
                continue
            for s in g["sales"]:
                self.sale_use_case.delete_sale(s.sale_id)
                total_borradas += 1
        MD.show_info(
            f"✅ {total_borradas} fiados eliminados y stock restaurado.",
            "Listo", parent=win)
        recargar()
        if self.current_page == "inventario":
            self.inventory_view.load_products()
        restaurar_foco()

    def on_ctrl_f12(event=None):
        tree = refs.get("tree")
        if not tree:
            return
        iid = None
        if marcados:
            iid = next(iter(marcados))
        else:
            sel = tree.selection()
            if sel:
                iid = sel[0]
        if not iid:
            MD.show_warning("Marca o selecciona un cliente primero.",
                            "Sin selección", parent=win)
            restaurar_foco()
            return
        g = grupos_map.get(iid)
        if not g or not g["sales"]:
            return
        venta_reciente = max(g["sales"], key=lambda x: x.sale_id)
        sid = venta_reciente.sale_id
        display = venta_reciente.display_number
        if MD.yesno(
                f"⚠️ ¿Eliminar el fiado más reciente de "
                f"'{g['customer_name']}'?\n\n"
                f"Venta #{display:02d} - ${venta_reciente.total:,.0f}\n"
                f"Esto restaurará el stock.".replace(",", "."),
                "Confirmar eliminación", parent=win) != "Yes":
            restaurar_foco()
            return
        pw = self._ask_password_1234(win)
        if not pw:
            restaurar_foco()
            return
        self.sale_use_case.delete_sale(sid)
        MD.show_info(
            f"Fiado #{display:02d} eliminado y stock restaurado.",
            "Listo", parent=win)
        recargar()
        if self.current_page == "inventario":
            self.inventory_view.load_products()
        restaurar_foco()

    def on_shift_f12(event=None):
        eliminar_marcados()

    def on_right_click(event):
        try:
            tree = refs.get("tree")
            row = tree.identify_row(event.y)
            if row:
                tree.selection_set(row)
                tree.focus(row)
            style = ttk.Style()
            m = tk.Menu(win, tearoff=0,
                        bg=style.colors.bg, fg=style.colors.fg,
                        activebackground=style.colors.selectbg,
                        activeforeground=style.colors.selectfg,
                        bd=1, relief="solid",
                        font=get_menu_font())
            m.add_command(label="👁️  Ver detalle", command=ver_detalle)
            m.add_command(label="💵 Abonar (más antigua)", command=abonar)
            m.add_command(label="✅ Marcar como pagado",
                          command=marcar_pagado)
            m.add_separator()
            if row:
                marca = "☑ Desmarcar" if row in marcados else "☐ Marcar"
                m.add_command(label=marca,
                              command=lambda r=row: toggle_mark(r))
            m.add_command(label="☑ Marcar todos", command=toggle_all)
            m.add_command(label="☐ Desmarcar todos",
                          command=desmarcar_todos)
            m.add_separator()
            m.add_command(
                label=f"🗑️  Eliminar marcados ({len(marcados)})",
                command=eliminar_marcados)
            try:
                m.tk_popup(event.x_root, event.y_root)
            finally:
                m.grab_release()
        except Exception:
            pass

    make_scrollable(container, build_content, build_bottom, bg=bg)
    recargar()

    show_popup_smooth(win)
    try:
        win.grab_set()
        win.focus_force()
    except Exception:
        pass

# ============================================================
# ABONAR
# ============================================================
def _abonar_dialog(self, parent, venta, pendiente, on_done):
    pop = tk.Toplevel(parent)
    pop.title(f"Abonar a venta #{venta.display_number:02d}")
    pop.geometry("440x520")
    pop.transient(parent)
    pop.configure(bg=self.style.colors.bg)
    pop.withdraw()
    bg = self.style.colors.bg
    fg = self.style.colors.fg

    container = tk.Frame(pop, bg=bg)
    container.pack(fill="both", expand=True)
    refs = {}

    def build_content(parent):
        tk.Label(parent, text="💵 Registrar Abono",
                 font=("Arial", 16, "bold"), bg=bg, fg=fg).pack(
                     pady=(15, 6))
        cliente = venta.customer_name if venta.customer_name else "(sin nombre)"
        tk.Label(parent, text=f"Cliente: {cliente}", font=("Arial", 12),
                 bg=bg, fg=fg).pack(pady=4)
        tk.Label(parent, text=f"Venta #{venta.display_number:02d}",
                 font=("Arial", 11, "italic"), bg=bg, fg="#a8e6a8").pack(
                     pady=2)
        tk.Label(parent,
                 text=f"Total: ${venta.total:,.0f}".replace(",", "."),
                 font=("Arial", 12), bg=bg, fg=fg).pack(pady=3)
        tk.Label(parent,
                 text=f"Abonado: ${venta.amount_paid:,.0f}".replace(
                     ",", "."),
                 font=("Arial", 12), bg=bg, fg=fg).pack(pady=3)
        tk.Label(parent,
                 text=f"Pendiente: ${pendiente:,.0f}".replace(",", "."),
                 font=("Arial", 14, "bold"),
                 bg="#0a4d1f", fg="#a8e6a8",
                 padx=10, pady=8).pack(pady=10)

        tk.Label(parent, text="Método del abono:",
                 font=("Arial", 11), bg=bg, fg=fg).pack(pady=(8, 3))
        metodo_var = tk.StringVar(value="Efectivo")
        mf = tk.Frame(parent, bg=bg)
        mf.pack(pady=3)
        for m in ["Efectivo", "Transferencia", "Tarjeta", "Otro"]:
            ttk.Radiobutton(mf, text=m, variable=metodo_var, value=m,
                            bootstyle="info").pack(side="left", padx=4)
        refs["metodo_var"] = metodo_var

        tk.Label(parent, text="Monto del abono:",
                 bg=bg, fg=fg).pack(pady=(10, 3))
        monto_var = tk.StringVar()
        e = ttk.Entry(parent, textvariable=monto_var, width=20,
                      font=("Arial", 16), justify="center")
        e.pack(pady=5)
        refs["entry"] = e
        refs["monto_var"] = monto_var

    def build_bottom(parent):
        def aplicar(ev=None):
            try:
                monto = float(refs["monto_var"].get().replace(
                    "$", "").replace(".", "").replace(",", "."))
                if monto <= 0:
                    raise ValueError
            except ValueError:
                MD.show_error("Monto inválido", "Error", parent=pop)
                return "break"
            if monto > pendiente + 0.01:
                if MD.yesno(
                        f"El monto (${monto:,.0f}) es mayor al pendiente "
                        f"(${pendiente:,.0f}).\n"
                        f"¿Registrar solo ${pendiente:,.0f}?".replace(
                            ",", "."),
                        "Confirmar", parent=pop) != "Yes":
                    return "break"
                monto = pendiente
            try:
                self.sale_use_case.add_payment(
                    venta.sale_id, monto,
                    method=refs["metodo_var"].get(),
                    customer_name=venta.customer_name or "")
                pop.destroy()
                MD.show_info(
                    f"✅ Abono de ${monto:,.0f} registrado.".replace(
                        ",", "."),
                    "Listo", parent=parent)
                on_done()
                try:
                    parent.lift()
                    parent.focus_force()
                except Exception:
                    pass
            except Exception as ex:
                MD.show_error(f"Error: {ex}", "Error", parent=pop)
            return "break"

        def cancelar(ev=None):
            pop.destroy()
            return "break"

        refs["aplicar"] = aplicar
        pop.bind("<Escape>", cancelar)

        ttk.Button(parent, text="Registrar abono", command=aplicar,
                   style="DarkGreen.TButton").pack(side="left", padx=5)
        ttk.Button(parent, text="Cancelar",
                   command=pop.destroy).pack(side="left", padx=5)

    make_scrollable(container, build_content, build_bottom, bg=bg)

    try:
        e = refs.get("entry")
        if e:
            e.bind("<Return>", refs.get("aplicar"))
            e.bind("<KP_Enter>", refs.get("aplicar"))
    except Exception:
        pass

    show_popup_smooth(pop)
    try:
        pop.grab_set()
        pop.focus_force()
    except Exception:
        pass

    def set_focus():
        try:
            if pop.winfo_exists():
                e = refs.get("entry")
                if e:
                    e.focus_set()
        except Exception:
            pass
    pop.after(50, set_focus)
    pop.after(250, set_focus)

# ============================================================
# AUTO-INICIO (Windows + Linux)
# ============================================================
def _get_startup_bat_path(self):
    startup = os.path.join(os.getenv('APPDATA'), 'Microsoft', 'Windows',
                           'Start Menu', 'Programs', 'Startup')
    return os.path.join(startup, 'MiniPOS_Portable.bat')

def _get_linux_autostart_path(self):
    autostart = os.path.expanduser("~/.config/autostart")
    return os.path.join(autostart, "MiniPOS.desktop")

def _is_autostart_enabled(self):
    try:
        if platform.system() == "Windows":
            return os.path.exists(self._get_startup_bat_path())
        elif platform.system() == "Linux":
            return os.path.exists(self._get_linux_autostart_path())
    except Exception:
        return False
    return False

def toggle_autostart(self):
    sistema = platform.system()
    if sistema == "Windows":
        bat = self._get_startup_bat_path()
        if self.autostart_var.get():
            try:
                with open(bat, 'w', encoding='utf-8') as f:
                    f.write(f'@echo off\nstart "" "{sys.executable}"\n')
                MD.show_info("✅ Auto-inicio ACTIVADO.", "Auto-inicio",
                             parent=self)
            except Exception as e:
                self.autostart_var.set(False)
                MD.show_error(f"Error: {e}", "Error", parent=self)
        else:
            try:
                if os.path.exists(bat):
                    os.remove(bat)
                MD.show_info("❌ Auto-inicio DESACTIVADO.",
                             "Auto-inicio", parent=self)
            except Exception as e:
                self.autostart_var.set(True)
                MD.show_error(f"Error: {e}", "Error", parent=self)
    elif sistema == "Linux":
        desktop = self._get_linux_autostart_path()
        if self.autostart_var.get():
            try:
                os.makedirs(os.path.dirname(desktop), exist_ok=True)
                if getattr(sys, 'frozen', False):
                    exe = sys.executable
                else:
                    exe = f"python3 {os.path.abspath(sys.argv[0])}"
                contenido = (
                    "[Desktop Entry]\n"
                    "Type=Application\n"
                    "Name=MiniPOS Portable\n"
                    "Comment=Punto de Venta\n"
                    f"Exec={exe}\n"
                    "Terminal=false\n"
                    "X-GNOME-Autostart-enabled=true\n"
                )
                with open(desktop, 'w', encoding='utf-8') as f:
                    f.write(contenido)
                os.chmod(desktop, 0o755)
                MD.show_info("✅ Auto-inicio ACTIVADO.", "Auto-inicio",
                             parent=self)
            except Exception as e:
                self.autostart_var.set(False)
                MD.show_error(f"Error: {e}", "Error", parent=self)
        else:
            try:
                if os.path.exists(desktop):
                    os.remove(desktop)
                MD.show_info("❌ Auto-inicio DESACTIVADO.",
                             "Auto-inicio", parent=self)
            except Exception as e:
                self.autostart_var.set(True)
                MD.show_error(f"Error: {e}", "Error", parent=self)
    else:
        MD.show_warning(f"Auto-inicio no soportado en {sistema}.",
                        "No soportado", parent=self)

# ============================================================
# REPORTES
# ============================================================
def generate_daily_report(self):
    hoy = datetime.now().strftime("%Y-%m-%d")
    sales = self.sale_use_case.get_sales_by_day(hoy)
    if not sales:
        MD.show_info("No hay ventas hoy.", "Reporte vacío", parent=self)
        return
    arch = filedialog.asksaveasfilename(
        defaultextension=".txt",
        initialfile=f"ventas_{hoy}.txt",
        filetypes=[("Texto", "*.txt")])
    if not arch:
        return
    total_dia = sum(s.total for s in sales)
    total_prod = sum(i.quantity for s in sales for i in s.items)
    L = ["=" * 60, "      REPORTE DE VENTAS DEL DÍA",
         f"      Fecha: {hoy}",
         f"      Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
         "=" * 60, ""]
    for s in sales:
        hora = s.date.split(" ")[1] if " " in s.date else ""
        cliente = f" - {s.customer_name}" if s.customer_name else ""
        fiado = " [FIADO]" if s.is_credit else ""
        L.append(f"Venta #{s.display_number:02d}  -  {hora}  -  "
                 f"{s.payment_method}{cliente}{fiado}")
        L.append(f"  {'Cant.':>8}  {'Producto':<30} "
                 f"{'P.Unit':>10} {'Subtotal':>12}")
        for it in s.items:
            L.append(f"  {it.quantity:>8g}  {it.product_name[:30]:<30} "
                     f"${it.unit_price:>9,.0f} "
                     f"${it.subtotal:>11,.0f}".replace(",", "."))
        if getattr(s, "discount", 0) > 0:
            L.append(f"  Descuento: -${s.discount:,.0f}".replace(",", "."))
        L.append(f"  {'':>8}  {'':<30} {'':>10} "
                 f"${s.total:>11,.0f}".replace(",", "."))
        L.append("")
    L += ["=" * 60,
          f"TOTAL DEL DÍA:      ${total_dia:>12,.0f}".replace(",", "."),
          f"VENTAS REALIZADAS:  {len(sales)}",
          f"PRODUCTOS VENDIDOS: {total_prod:g}", "=" * 60]
    try:
        with open(arch, "w", encoding="utf-8") as f:
            f.write("\n".join(L))
        MD.show_info(f"Guardado:\n{arch}", "Listo", parent=self)
    except Exception as e:
        MD.show_error(f"Error: {e}", "Error", parent=self)

def generate_monthly_report(self):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        MD.show_error("Falta openpyxl.", "Error", parent=self)
        return
    ym = datetime.now().strftime("%Y-%m")
    sales = self.sale_use_case.get_sales_by_month(ym)
    if not sales:
        MD.show_info("No hay ventas este mes.", "Reporte vacío",
                     parent=self)
        return
    arch = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        initialfile=f"ventas_{ym}.xlsx",
        filetypes=[("Excel", "*.xlsx")])
    if not arch:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Ventas {ym}"
    ws.append(["Venta #", "Fecha", "Cliente", "Método", "Fiado",
               "Producto", "Código", "Cantidad", "Unidad",
               "P. Unit.", "Subtotal Item",
               "Desc. Venta", "Total Venta"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    for s in sales:
        for i, it in enumerate(s.items):
            desc_venta = getattr(s, "discount", 0) if i == 0 else ""
            total_venta = s.total if i == 0 else ""
            ws.append([
                f"#{s.display_number:02d}", s.date, s.customer_name,
                s.payment_method, "Sí" if s.is_credit else "No",
                it.product_name, it.barcode,
                it.quantity, "unidad", it.unit_price, it.subtotal,
                desc_venta, total_venta
            ])
    total = sum(s.total for s in sales)
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "", "", "",
               "TOTAL:", total])
    for col, w in zip("ABCDEFGHIJKLM",
                      [10, 20, 25, 15, 8, 30, 20, 10, 10,
                       12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    try:
        wb.save(arch)
        MD.show_info(f"Guardado:\n{arch}", "Listo", parent=self)
    except Exception as e:
        MD.show_error(f"Error: {e}", "Error", parent=self)

# ============================================================
# EXPORTAR / IMPORTAR BD
# ============================================================
def export_db(self):
    arch = filedialog.asksaveasfilename(
        defaultextension=".db",
        filetypes=[("SQLite", "*.db")])
    if arch:
        self.db_manager.close_connection()
        shutil.copy2(self.db_path, arch)
        MD.show_info("Base de datos exportada.", "Listo", parent=self)

def import_db(self):
    arch = filedialog.askopenfilename(filetypes=[("SQLite", "*.db")])
    if arch and MD.yesno("¿Reemplazar datos?", "Confirmar",
                         parent=self) == "Yes":
        self.db_manager.close_connection()
        shutil.copy2(arch, self.db_path)
        self.inventory_view.load_products()
        MD.show_info("Importada.", "Listo", parent=self)
